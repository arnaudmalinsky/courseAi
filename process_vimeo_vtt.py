"""Nettoie et synthétise en lot les transcriptions Vimeo au format WebVTT."""

from __future__ import annotations

import argparse
import asyncio
import json
import re
import time
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

COMMON_PROMPT = """Travaille exclusivement à partir de la retranscription fournie. N'ajoute aucune jurisprudence, aucun article, aucune règle et aucune explication provenant de connaissances extérieures.

La fidélité et la complétude priment sur la concision. Ne résume pas globalement la conférence. Procède essentiellement par suppression des répétitions strictement identiques, hésitations, tics de langage, phrases inachevées sans contenu et éléments expressément exclus par les instructions propres au type de document. Conserve tous les développements juridiques distincts, même lorsqu'ils paraissent secondaires.

Ne supprime aucune règle, condition, exception, distinction, qualification, hypothèse, réserve, nuance, exemple juridique, étape de raisonnement ou conseil méthodologique développé par le professeur. Ne fusionne pas des développements voisins lorsqu'ils apportent des informations différentes.

Restitue uniquement le texte final, sans préambule ni commentaire sur ton travail. Corrige une erreur manifeste de transcription seulement lorsque le contexte interne permet de le faire sans ambiguïté.

Organise le contenu en paragraphes lisibles et relativement courts, sans réduire la substance juridique. N'utilise aucune liste à puces, aucune liste numérotée, aucun tableau, aucun emoji et aucun séparateur horizontal. Utilise des titres et sous-titres descriptifs en Markdown (#, ## ou ###) afin qu'ils soient convertis ensuite en styles Word natifs. Sous chaque titre, écris uniquement des paragraphes rédigés.

N'utilise jamais « Idée principale », « Jurisprudence », « Portée » ou « Norme juridique » comme titre ou catégorie séparée.

Conserve toute référence juridique citée, même brièvement. Ne supprime et ne généralise jamais le numéro d'un article, le nom d'une décision, sa juridiction ou sa date lorsqu'ils figurent dans la retranscription. Lorsque tous ces éléments sont disponibles, cite une jurisprudence selon le formalisme « juridiction, année, nom », par exemple « CE, 1933, Benjamin ». N'invente jamais un élément manquant pour compléter ce formalisme.

ATTENTION AUX RÉFÉRENCES JURIDIQUES : une référence juridique peut prendre des formes très diverses dans une conférence orale. Elle peut notamment être annoncée par les mots « loi », « article », « code », « décret », « ordonnance », « règlement », « directive », « jurisprudence », « arrêt », « arrêt du Conseil d'État », « décision du Conseil constitutionnel », « décision de la Cour de cassation », « décision du Tribunal des conflits », « décision de... », « vous citez », « vous visez » ou « vous vous fondez sur ». Cette liste n'est pas exhaustive. Considère toute référence de ce type comme une information prioritaire à conserver, ainsi que l'explication et le raisonnement auxquels le professeur la rattache.

Le formalisme de citation doit s'adapter aux informations réellement présentes dans la retranscription et n'est donc pas entièrement figé. Pour une jurisprudence, privilégie l'ordre « juridiction, date ou année, nom de la décision » lorsque ces éléments sont prononcés, par exemple « CE, 1933, Benjamin » ou « CE, 1er avril 1994, Commune de Menton ». Pour une loi, conserve sa date et son intitulé ou surnom lorsqu'ils sont donnés. Pour un article, conserve son numéro et le code auquel il appartient lorsqu'il est cité. Pour une décision constitutionnelle, judiciaire ou européenne, conserve la juridiction, la date, le numéro et le nom uniquement dans la mesure où ils figurent dans la source. Ne complète jamais une citation à l'aide de connaissances extérieures.

Une référence reste obligatoire même si elle est répartie sur plusieurs phrases, si son nom apparaît après une coupure de la retranscription ou si le professeur ne la mentionne qu'une seule fois. Les expressions « vous citez », « vous visez », « vous vous fondez sur » et les formulations équivalentes signalent généralement une référence à reprendre avec une attention particulière.

"""

ACTU_PROMPT = """Reprends cette retranscription d'une conférence d'actualité en améliorant la mise en page et en supprimant les apartés.

Structure le document autour de chaque jurisprudence commentée. Pour chacune, reprends avec précision tout le contenu développé par le professeur : les faits mentionnés, la question examinée, la solution, le raisonnement, les règles mobilisées, les nuances, les conséquences et les rapprochements avec d'autres décisions, uniquement lorsqu'ils apparaissent dans la retranscription.

Pour chaque jurisprudence, conserve chaque information distincte effectivement donnée par le professeur concernant le contexte, les faits, la procédure, la question juridique, la solution, la motivation, la portée expliquée, les limites, les critiques, les évolutions et les comparaisons. Si un élément n'est pas abordé, ne l'invente pas. S'il est abordé, ne l'omets pas.

Ne réduis jamais une jurisprudence à son seul nom ou à une phrase générale. Si le professeur revient ultérieurement sur une décision ou complète son commentaire, réintègre ces développements dans la section correspondante.
"""

CORR_PROMPT = """Reprends cette retranscription d'une séance de correction de cas pratique en améliorant la mise en page et en reprenant avec précision, sans ajout extérieur, la réponse développée pour chaque cas pratique.

Conserve les apartés, mais place chaque aparté dans un paragraphe distinct clairement intitulé « Aparté », sans le mélanger au raisonnement juridique.

Pour chaque cas pratique, distingue clairement chaque question traitée. Pour chaque question, présente successivement des sous-titres « Question de droit », « Majeure », « Mineure » et « Conclusion ».

La question de droit doit reprendre le problème juridique effectivement identifié pendant la correction. La majeure doit conserver toutes les règles, tous les articles et toutes les jurisprudences mobilisés. La mineure doit reprendre précisément l'application aux faits. La conclusion doit restituer la solution retenue ou envisagée. Si une étape est développée plus tard dans la séance, réintègre-la dans la question correspondante.

Ne fusionne jamais plusieurs questions de droit. Ne réduis pas un syllogisme à quelques phrases lorsque la correction est plus développée. Conserve toutes les hypothèses, objections, qualifications alternatives, moyens, conditions, exceptions, références, articulations du raisonnement et précautions méthodologiques présentés par le professeur.
"""

COURSE_PROMPT = """Reprends cette retranscription d'un cours juridique en améliorant la mise en page, en supprimant les apartés et en conservant avec précision les règles, raisonnements, articles, jurisprudences, exemples et nuances développés par le professeur.
"""


def prompt_for_source(source: Path) -> str:
    stem = source.stem.lower()
    if "_actu_" in stem:
        specific = ACTU_PROMPT
    elif "_corr_" in stem:
        specific = CORR_PROMPT
    else:
        specific = COURSE_PROMPT
    return f"{specific}\n{COMMON_PROMPT}\nRETRANSCRIPTION À TRAITER :\n"

TIMESTAMP_RE = re.compile(
    r"^\s*\d{2}:\d{2}(?::\d{2})?[.,]\d{3}\s+-->\s+"
    r"\d{2}:\d{2}(?::\d{2})?[.,]\d{3}.*$"
)
FORBIDDEN_RE = re.compile(
    r"\b(Idée principale|Jurisprudence\s*:|Portée\s*:|Norme juridique)\b",
    re.IGNORECASE,
)


def log_event(log_path: Path, event: str, **details) -> None:
    record = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "event": event,
        **details,
    }
    line = json.dumps(record, ensure_ascii=False)
    print(line, flush=True)
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(line + "\n")


def load_settings(env_path: str) -> dict:
    values: dict[str, str] = {}
    for raw in Path(env_path).read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            values[key.strip()] = value.strip().strip("\"'")
    effort = values.get("MODEL_REASONING") or values.get("MODEL_REASONNING")
    if effort and effort.lower() == "standard":
        effort = "medium"
    return {
        "api_key": values.get("OPENAI_API_KEY") or values.get("OPENAI_KEY"),
        "model": values.get("OPENAI_MODEL") or "gpt-4o",
        "reasoning_effort": effort,
        "verbosity": values.get("MODEL_VERBOSITY"),
        "reasoning_summary": values.get("MODEL_SUMMARY"),
    }


def call_openai_stream_sync(
    settings: dict,
    prompt: str,
    partial_path: Path,
    log_path: Path,
    source_name: str,
) -> str:
    payload: dict = {"model": settings["model"], "input": prompt, "stream": True}
    reasoning = {}
    if settings.get("reasoning_effort"):
        reasoning["effort"] = settings["reasoning_effort"]
    if settings.get("reasoning_summary"):
        reasoning["summary"] = settings["reasoning_summary"]
    if reasoning:
        payload["reasoning"] = reasoning
    if settings.get("verbosity"):
        payload["text"] = {"verbosity": settings["verbosity"]}
    request = urllib.request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {settings['api_key']}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        chunks: list[str] = []
        received_chars = 0
        last_progress = time.monotonic()
        last_reported_chars = 0
        completed = False
        with partial_path.open("w", encoding="utf-8", newline="") as partial:
            with urllib.request.urlopen(request, timeout=900) as response:
                for raw_line in response:
                    line = raw_line.decode("utf-8", errors="replace").strip()
                    if not line.startswith("data:"):
                        continue
                    data_text = line[5:].strip()
                    if not data_text or data_text == "[DONE]":
                        continue
                    event = json.loads(data_text)
                    event_type = event.get("type", "")
                    if event_type == "response.output_text.delta":
                        delta = event.get("delta", "")
                        if delta:
                            chunks.append(delta)
                            partial.write(delta)
                            partial.flush()
                            received_chars += len(delta)
                    elif event_type == "response.completed":
                        completed = True
                    elif event_type in {"error", "response.failed", "response.incomplete"}:
                        error = event.get("error") or event.get("response", {}).get("error") or event
                        raise RuntimeError(f"Échec du flux OpenAI: {error}")

                    now = time.monotonic()
                    if (
                        received_chars - last_reported_chars >= 1000
                        or now - last_progress >= 30
                    ):
                        log_event(
                            log_path,
                            "stream_progress",
                            file=source_name,
                            received_chars=received_chars,
                            partial_file=str(partial_path),
                        )
                        last_reported_chars = received_chars
                        last_progress = now
        result = "".join(chunks)
        if not completed:
            raise RuntimeError("Le flux OpenAI s'est fermé sans événement response.completed")
        if not result.strip():
            raise ValueError("Le flux OpenAI ne contient aucun texte exploitable")
        return result
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")[:2000]
        raise RuntimeError(f"HTTP {exc.code}: {detail}") from exc


async def call_openai(
    settings: dict,
    prompt: str,
    partial_path: Path,
    log_path: Path,
    source_name: str,
) -> str:
    return await asyncio.to_thread(
        call_openai_stream_sync,
        settings,
        prompt,
        partial_path,
        log_path,
        source_name,
    )


def vtt_to_text(path: Path) -> str:
    lines: list[str] = []
    for raw in path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        line = raw.strip()
        if not line or line == "WEBVTT" or line.isdigit() or TIMESTAMP_RE.match(line):
            continue
        if line.startswith(("NOTE", "STYLE", "REGION")):
            continue
        line = re.sub(r"<[^>]+>", "", line)
        lines.append(line)
    text = " ".join(lines)
    return re.sub(r"\s+", " ", text).strip()


def validate(text: str, source_length: int | None = None) -> list[str]:
    errors: list[str] = []
    if not text.strip():
        errors.append("résultat vide")
    if TIMESTAMP_RE.search(text):
        errors.append("minutage restant")
    if FORBIDDEN_RE.search(text):
        errors.append("catégorie interdite")
    if re.search(r"(?m)^\s*\|.*\|\s*$", text):
        errors.append("tableau Markdown détecté")
    if re.search(r"(?m)^\s*(?:---+|___+|\*\*\*+)\s*$", text):
        errors.append("séparateur détecté")
    if re.search(r"(?m)^\s*(?:[-*+]\s+|\d+[.)]\s+)", text):
        errors.append("liste à puces ou numérotée détectée")
    return errors


def document_title(stem: str) -> str:
    labels = {"DA": "Droit administratif", "DO": "Droit des obligations", "PAC": "Contentieux administratif"}
    kinds = {"corr": "Correction", "actu": "Actualité", "cours": "Cours"}
    parts = stem.split("_")
    subject = labels.get(parts[0], parts[0])
    kind = kinds.get(parts[1], parts[1].capitalize()) if len(parts) > 1 else ""
    number = parts[2] if len(parts) > 2 else ""
    return " - ".join(p for p in (subject, f"{kind} {number}".strip()) if p)


def set_font(style, name: str, size: float, bold: bool = False, color: str = "000000") -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))


def write_docx(markdown: str, output: Path, title: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_font(normal, "Calibri", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        set_font(style, "Calibri", size, True, color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(14)
    title_p.paragraph_format.keep_with_next = True
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor.from_string("0B2545")
    title_run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    title_run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")

    for block in re.split(r"\n\s*\n", markdown.strip()):
        block = block.strip()
        if not block:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", block)
        if heading and "\n" not in block:
            heading_level = min(len(heading.group(1)), 3)
            doc.add_paragraph(heading.group(2).strip(), style=f"Heading {heading_level}")
            continue
        for line in block.splitlines():
            line = line.strip()
            if not line:
                continue
            bullet = re.match(r"^[-*+]\s+(.+)$", line)
            numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
            if bullet:
                doc.add_paragraph(bullet.group(1), style="List Bullet")
            elif numbered:
                doc.add_paragraph(numbered.group(1), style="List Number")
            else:
                doc.add_paragraph(line)

    add_page_number(section.footer.paragraphs[0])
    doc.core_properties.title = title
    doc.core_properties.author = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


async def process_one(
    source: Path,
    output_dir: Path,
    settings: dict,
    semaphore: asyncio.Semaphore,
    force: bool,
    log_path: Path,
) -> dict:
    output = output_dir / f"{source.stem}.md"
    partial_output = output_dir / f"{source.stem}.partial.md"
    docx_output = output_dir / f"{source.stem}.docx"
    llm_input_output = output_dir / f"{source.stem}.llm_input.txt"
    if output.exists() and output.stat().st_size > 100 and not force:
        if not docx_output.exists():
            write_docx(output.read_text(encoding="utf-8"), docx_output, document_title(source.stem))
        return {"file": source.name, "status": "skipped", "output": str(docx_output)}

    transcript = vtt_to_text(source)
    llm_input = prompt_for_source(source) + transcript
    llm_input_output.write_text(llm_input, encoding="utf-8")
    try:
        log_event(
            log_path,
            "api_call_started",
            file=source.name,
            model=settings["model"],
            source_chars=len(transcript),
            partial_file=str(partial_output),
            llm_input_file=str(llm_input_output),
            attempt=1,
        )
        started = time.monotonic()
        async with semaphore:
            api_task = asyncio.create_task(
                call_openai(
                    settings,
                    llm_input,
                    partial_output,
                    log_path,
                    source.name,
                )
            )
            while True:
                done, _ = await asyncio.wait({api_task}, timeout=30)
                if done:
                    result = api_task.result()
                    break
                log_event(
                    log_path,
                    "api_call_waiting",
                    file=source.name,
                    elapsed_seconds=round(time.monotonic() - started, 1),
                )
        warnings = validate(result)
        log_event(
            log_path,
            "api_call_completed",
            file=source.name,
            elapsed_seconds=round(time.monotonic() - started, 1),
            output_chars=len(result.strip()),
            validation_warnings=warnings,
        )
        with partial_output.open("a", encoding="utf-8") as partial:
            partial.write("\n")
        partial_output.replace(output)
        write_docx(result, docx_output, document_title(source.stem))
        return {
            "file": source.name,
            "status": "completed",
            "output": str(docx_output),
            "markdown_output": str(output),
            "llm_input": str(llm_input_output),
            "source_chars": len(transcript),
            "output_chars": len(result.strip()),
            "validation_warnings": warnings,
            "attempts": 1,
        }
    except Exception as exc:
        log_event(
            log_path,
            "fatal_error",
            file=source.name,
            error_type=type(exc).__name__,
            error=str(exc),
        )
        return {
            "file": source.name,
            "status": "error",
            "error": f"{type(exc).__name__}: {exc}",
            "source_chars": len(transcript),
            "attempts": 1,
        }


async def run(args: argparse.Namespace) -> int:
    source_dir = Path(args.input).resolve()
    output_dir = Path(args.output).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    log_path = output_dir / "run.jsonl"
    sources = sorted(source_dir.glob("*.vtt"), key=lambda p: p.name.lower())
    if args.file:
        sources = [p for p in sources if p.name.lower() == args.file.lower()]
    if args.limit:
        sources = sources[: args.limit]
    if not sources:
        raise SystemExit(f"Aucun fichier VTT trouvé dans {source_dir}")

    settings = load_settings(args.env)
    if not settings.get("api_key"):
        raise SystemExit("Clé OpenAI absente du fichier .env")
    log_event(
        log_path,
        "run_started",
        model=settings["model"],
        file_count=len(sources),
        concurrency=args.concurrency,
    )
    semaphore = asyncio.Semaphore(args.concurrency)
    tasks = [
        asyncio.create_task(
            process_one(s, output_dir, settings, semaphore, args.force, log_path)
        )
        for s in sources
    ]
    results: list[dict] = []
    for task in asyncio.as_completed(tasks):
        item = await task
        results.append(item)
        manifest = {
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "model": settings["model"],
            "input_dir": str(source_dir),
            "output_dir": str(output_dir),
            "results": sorted(results, key=lambda x: x["file"].lower()),
        }
        (output_dir / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        if item["status"] == "error":
            log_event(
                log_path,
                "run_aborted",
                file=item["file"],
                error=item.get("error", "Erreur inconnue"),
            )
            for pending in tasks:
                if not pending.done():
                    pending.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
            return 1
        log_event(log_path, "file_completed", file=item["file"])
    log_event(log_path, "run_completed", file_count=len(results))
    return 0


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", default="data/vimeo_processed")
    parser.add_argument("--env", default=".env")
    parser.add_argument("--concurrency", type=int, default=1)
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--file")
    parser.add_argument("--limit", type=int)
    raise SystemExit(asyncio.run(run(parser.parse_args())))


if __name__ == "__main__":
    main()
