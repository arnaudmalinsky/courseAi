import logging
import re
from pathlib import Path
from statistics import median

from docx import Document


class PdfBlock:
    def __init__(self, page_number, text, font_size, is_bold):
        self.page_number = page_number
        self.text = text
        self.font_size = font_size
        self.is_bold = is_bold


def convert_pdf_to_docx(
    input_pdf_path,
    output_docx_path=None,
    preserve_page_breaks=False,
    max_heading_chars=140,
    chapters_dir=None,
    manifest_path=None,
):
    input_path = Path(input_pdf_path)
    output_path = Path(output_docx_path) if output_docx_path else input_path.with_suffix(".docx")

    if input_path.suffix.lower() != ".pdf":
        raise ValueError(f"Input file must be a PDF: {input_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        import pymupdf
    except ImportError as exc:
        raise ImportError(
            "PyMuPDF is required for PDF conversion. Install it with: pip install pymupdf"
        ) from exc

    with pymupdf.open(input_path) as pdf_document:
        blocks = _extract_blocks(pdf_document)

    if not blocks:
        raise ValueError(f"No selectable text was found in PDF: {input_path}")

    body_font_size = _get_body_font_size(blocks)
    heading_size_levels = _get_heading_size_levels(blocks, body_font_size)

    logging.info("Body font size detected: %s", body_font_size)
    logging.info("Heading font sizes detected: %s", heading_size_levels)

    doc = Document()

    previous_page_number = blocks[0].page_number
    for block in blocks:
        if preserve_page_breaks and block.page_number != previous_page_number:
            doc.add_page_break()
            previous_page_number = block.page_number

        heading_level = _detect_heading_level(
            block,
            body_font_size,
            heading_size_levels,
            max_heading_chars,
        )

        if heading_level:
            doc.add_heading(block.text, level=heading_level)
        else:
            doc.add_paragraph(block.text)

    doc.save(output_path)
    logging.warning("DOCX saved to %s", output_path)

    if chapters_dir:
        from .chapter_pipeline import split_docx_by_chapter

        split_docx_by_chapter(output_path, chapters_dir, manifest_path)

    return output_path


def _extract_blocks(pdf_document):
    blocks = []
    for page_index, page in enumerate(pdf_document):
        page_dict = page.get_text("dict")

        for block in page_dict.get("blocks", []):
            if block.get("type") != 0:
                continue

            block_data = _parse_text_block(page_index + 1, block)
            if block_data is not None:
                blocks.append(block_data)

    return blocks


def _parse_text_block(page_number, block):
    lines = []
    font_sizes = []
    bold_votes = []

    for line in block.get("lines", []):
        line_parts = []
        for span in line.get("spans", []):
            text = span.get("text", "")
            if text.strip():
                line_parts.append(text.strip())
                font_sizes.append(round(float(span.get("size", 0)), 1))
                bold_votes.append(bool(span.get("flags", 0) & 16))

        if line_parts:
            lines.append(" ".join(line_parts))

    text = _clean_block_text(lines)
    if not text:
        return None

    font_size = median(font_sizes) if font_sizes else 0
    is_bold = bool(bold_votes) and sum(bold_votes) >= len(bold_votes) / 2

    return PdfBlock(page_number, text, font_size, is_bold)


def _clean_block_text(lines):
    cleaned_lines = [line.strip() for line in lines if line.strip()]
    if not cleaned_lines:
        return ""

    text = " ".join(cleaned_lines)
    text = _normalize_pdf_heading_spacing(text)
    return " ".join(text.split())


def _normalize_pdf_heading_spacing(text):
    replacements = {
        "S ECTION": "SECTION",
        "C HAPITRE": "CHAPITRE",
        "T HÈME": "THÈME",
        "T HEME": "THEME",
        "T ITRE": "TITRE",
        "P ARTIE": "PARTIE",
        "L ES": "LES",
        "L A": "LA",
        "L E": "LE",
        "C ONSTITUTION": "CONSTITUTION",
        "U NION": "UNION",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _get_body_font_size(blocks):
    font_sizes = [block.font_size for block in blocks if block.font_size > 0]
    if not font_sizes:
        return 0
    return median(font_sizes)


def _get_heading_size_levels(blocks, body_font_size):
    heading_sizes = sorted(
        {
            block.font_size
            for block in blocks
            if block.font_size >= body_font_size + 1.2
        },
        reverse=True,
    )

    return {
        font_size: level
        for level, font_size in enumerate(heading_sizes[:3], start=1)
    }


def _detect_heading_level(
    block,
    body_font_size,
    heading_size_levels,
    max_heading_chars,
):
    if len(block.text) > max_heading_chars:
        return None

    plan_heading_level = _get_plan_heading_level(block.text)
    if plan_heading_level is not None:
        return plan_heading_level

    if block.font_size in heading_size_levels:
        return heading_size_levels[block.font_size]

    if block.is_bold and block.font_size >= body_font_size + 0.8:
        return 3

    return None


def _get_plan_heading_level(text):
    normalized_text = text.strip().upper()
    normalized_text = normalized_text.replace("È", "E")

    if re.match(r"^PARTIE\s+[IVXLCDM]+\b", normalized_text):
        return 1

    if re.match(r"^THEME\s+[IVXLCDM]+\b", normalized_text):
        return 2

    if re.match(r"^CHAPITRE\s+([IVXLCDM]+|\d+)\b", normalized_text):
        return 3

    if re.match(r"^SECTION\s+([IVXLCDM]+|\d+)\b", normalized_text):
        return 4

    if re.match(r"^(I|II|III|IV|V|VI|VII|VIII|IX|X)\.\s+\S", normalized_text):
        return 5

    if re.match(r"^[A-Z]\.\s+(?![A-Z]\.)\S", normalized_text):
        return 6

    if re.match(r"^\d+\)\s+\S", normalized_text):
        return 7

    return None
