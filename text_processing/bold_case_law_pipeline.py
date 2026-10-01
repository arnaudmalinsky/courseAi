import asyncio
import base64
import json
import logging
import re
from dataclasses import asdict, dataclass
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from .chapter_pipeline import _call_openai_responses, _resolve_openai_settings


BOLD_CASE_LAW_PROMPT = """À partir de ce chapitre de manuel PDF, dresse un répertoire exhaustif des jurisprudences dont la référence est typographiée en gras.

Consignes impératives :
- Ne retiens que les jurisprudences dont la référence juridique elle-même apparaît en gras dans le PDF.
- N'ajoute jamais une jurisprudence dont la référence apparaît uniquement en caractères ordinaires.
- N'ajoute aucune jurisprudence issue de tes connaissances personnelles.
- Utilise exclusivement le contexte du manuel pour synthétiser la portée juridique.
- Respecte strictement le plan du manuel.
- Reprends exactement les intitulés des Parties, Thèmes, Titres, Chapitres, Sections et Sous-sections.
- Classe chaque jurisprudence dans la partie du plan où elle apparaît.
- Lorsque plusieurs références en gras sont regroupées par le manuel et partagent exactement la même portée juridique, conserve-les ensemble dans un seul groupe.
- Ne regroupe pas des références seulement parce qu'elles sont proches : elles doivent être explicitement associées à la même portée dans le manuel.
- Dans chaque groupe, conserve l'ordre d'apparition des références.
"""

BOLD_TECHNICAL_INSTRUCTIONS = """
INSTRUCTIONS TECHNIQUES DE SORTIE :
- Analyse visuellement toutes les pages du chapitre afin d'identifier la graisse typographique.
- Le champ « references » contient une ou plusieurs références juridiques, sans portée ni deux-points final.
- Le champ « description » contient uniquement la portée juridique commune, sans répéter les références.
- Une référence isolée constitue un groupe d'un seul élément.
- Le champ « source_page » correspond au numéro de page du PDF source indiqué dans les métadonnées de découpage ci-dessous.
- N'invente aucun niveau de plan absent du PDF. Laisse la chaîne vide si un niveau n'est pas visible.
- Reproduis les intitulés du plan caractère pour caractère, sans les corriger ni les reformuler.
- N'ajoute aucun commentaire, avertissement ou texte hors du format structuré demandé.
"""


@dataclass(frozen=True)
class PdfPosition:
    page_index: int
    y: float


@dataclass(frozen=True)
class Heading:
    kind: str
    title: str
    page_index: int
    y: float


@dataclass(frozen=True)
class ChapterSlice:
    order: int
    title: str
    start_page_index: int
    start_y: float
    end_page_index: int
    end_y: float


class GroupedCaseLawItem(BaseModel):
    model_config = ConfigDict(extra="forbid")

    references: list[str] = Field(
        min_length=1,
        description="Références en gras partageant exactement la même portée",
    )
    description: str = Field(description="Portée juridique commune")
    source_page: int | None = Field(description="Page du PDF source original")

    @property
    def reference_text(self):
        return " ; ".join(ref.strip().rstrip(" :") for ref in self.references)

    @property
    def formulation(self):
        return f"{self.reference_text} : {self.description.strip().lstrip(' :')}"


class BoldPlanEntry(BaseModel):
    model_config = ConfigDict(extra="forbid")

    partie: str
    theme: str
    titre: str
    chapitre: str
    section: str
    sous_section: str
    groupes: list[GroupedCaseLawItem]


class BoldChapterExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")

    entries: list[BoldPlanEntry]


def _normalize_heading_text(text):
    text = re.sub(r"\s+", " ", text).strip()
    replacements = {
        "C HAPITRE": "CHAPITRE",
        "P ARTIE": "PARTIE",
        "T HÈME": "THÈME",
        "T HEME": "THEME",
        "T ITRE": "TITRE",
        "L ES": "LES",
        "L A": "LA",
        "L E": "LE",
        "D ’": "D’",
        "D '": "D'",
        "S ERVICE": "SERVICE",
        "R ÉTENTION": "RÉTENTION",
        "A UTRES": "AUTRES",
        "C ONTENTIEUX": "CONTENTIEUX",
    }
    for source, target in replacements.items():
        text = re.sub(rf"\b{re.escape(source)}", target, text, flags=re.IGNORECASE)
    return text


def _heading_kind(text):
    normalized = _normalize_heading_text(text).upper()
    patterns = (
        ("partie", r"^PARTIE\s+(?:[IVXLCDM]+|\d+)\b"),
        ("theme", r"^TH[ÈE]ME\s+(?:[IVXLCDM]+|\d+)\b"),
        ("titre", r"^TITRE(?:\s+(?:[IVXLCDM]+|\d+)|\s+LIMINAIRE)\b"),
        ("chapitre", r"^CHAPITRE\s+(?:[IVXLCDM]+|\d+)\b"),
    )
    for kind, pattern in patterns:
        if re.match(pattern, normalized):
            return kind
    return None


def _detect_course_end(pdf_document):
    for page_index, page in enumerate(pdf_document):
        if page_index < 250:
            continue
        text = page.get_text("text")
        chapter_toc_lines = sum(
            bool(re.search(r"Chapitre\s+[IVXLCDM\d]+.*\.{5,}\s*\d+", line, re.I))
            for line in text.splitlines()
        )
        if chapter_toc_lines >= 2:
            return PdfPosition(page_index, 0.0)
    return PdfPosition(pdf_document.page_count, 0.0)


def detect_pdf_headings(pdf_document):
    headings = []
    course_end = _detect_course_end(pdf_document)
    for page_index in range(course_end.page_index):
        page = pdf_document[page_index]
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            spans = [
                span
                for line in block.get("lines", [])
                for span in line.get("spans", [])
                if span.get("text", "").strip()
            ]
            if not spans or max(float(span.get("size", 0)) for span in spans) < 14:
                continue
            text = _normalize_heading_text(" ".join(span["text"].strip() for span in spans))
            kind = _heading_kind(text)
            if kind:
                headings.append(
                    Heading(kind, text, page_index, round(float(block["bbox"][1]), 2))
                )
    return headings, course_end


def build_chapter_slices(pdf_document):
    headings, course_end = detect_pdf_headings(pdf_document)
    chapters = [heading for heading in headings if heading.kind == "chapitre"]
    if not chapters:
        raise ValueError("Aucun titre de chapitre n'a été détecté dans le PDF.")

    boundaries = []
    for chapter_index, chapter in enumerate(chapters):
        previous = chapters[chapter_index - 1] if chapter_index else None
        candidates = [
            heading
            for heading in headings
            if heading.kind in {"partie", "theme", "titre"}
            and (heading.page_index, heading.y) <= (chapter.page_index, chapter.y)
            and (
                previous is None
                or (heading.page_index, heading.y) > (previous.page_index, previous.y)
            )
        ]
        start_heading = candidates[0] if candidates else chapter
        boundaries.append(PdfPosition(start_heading.page_index, start_heading.y))

    slices = []
    for index, (chapter, start) in enumerate(zip(chapters, boundaries), start=1):
        end = boundaries[index] if index < len(boundaries) else course_end
        if (start.page_index, start.y) >= (end.page_index, end.y):
            raise ValueError(f"Frontière invalide pour {chapter.title}")
        slices.append(
            ChapterSlice(
                order=index,
                title=chapter.title,
                start_page_index=start.page_index,
                start_y=start.y,
                end_page_index=end.page_index,
                end_y=end.y,
            )
        )
    return slices, course_end


def _safe_chapter_filename(chapter_slice):
    slug = re.sub(r"[^\w.-]+", "_", chapter_slice.title, flags=re.UNICODE).strip("_.")
    return f"chapitre_{chapter_slice.order:03d}_{slug[:80]}.pdf"


def _append_clipped_page(target, source, page_index, y0, y1):
    import fitz

    source_page = source[page_index]
    page_rect = source_page.rect
    y0 = max(page_rect.y0, float(y0))
    y1 = min(page_rect.y1, float(y1))
    if y1 - y0 < 1:
        return
    clip = fitz.Rect(page_rect.x0, y0, page_rect.x1, y1)
    output_page = target.new_page(width=clip.width, height=clip.height)
    # Rasterisation volontaire : show_pdf_page/cropbox peut conserver du texte
    # masqué hors de la zone visible, que l'API lirait malgré le rognage.
    pixmap = source_page.get_pixmap(matrix=fitz.Matrix(2, 2), clip=clip, alpha=False)
    output_page.insert_image(output_page.rect, stream=pixmap.tobytes("png"))


def _write_chapter_pdf(source, chapter_slice, output_path):
    import fitz

    target = fitz.open()
    start_page = chapter_slice.start_page_index
    end_page = chapter_slice.end_page_index
    if start_page == end_page:
        _append_clipped_page(
            target, source, start_page, chapter_slice.start_y, chapter_slice.end_y
        )
    else:
        _append_clipped_page(
            target, source, start_page, chapter_slice.start_y, source[start_page].rect.y1
        )
        for page_index in range(start_page + 1, end_page):
            _append_clipped_page(
                target, source, page_index, source[page_index].rect.y0, source[page_index].rect.y1
            )
        # Une frontière vers y=74 correspond au titre situé tout en haut de la
        # page suivante : la bande qui précède ne contient que l'en-tête courant.
        if end_page < source.page_count and chapter_slice.end_y > 100:
            _append_clipped_page(
                target, source, end_page, source[end_page].rect.y0, chapter_slice.end_y
            )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.is_file():
        output_path.unlink()
    target.save(output_path, garbage=4, deflate=True)
    target.close()


def split_pdf_by_detected_chapters(input_pdf_path, output_dir, manifest_path=None):
    import fitz

    input_path = Path(input_pdf_path).resolve()
    output_path = Path(output_dir).resolve()
    manifest = (
        Path(manifest_path).resolve()
        if manifest_path
        else output_path / "chapters_manifest.json"
    )
    with fitz.open(input_path) as source:
        slices, course_end = build_chapter_slices(source)
        records = []
        for chapter_slice in slices:
            chapter_path = output_path / _safe_chapter_filename(chapter_slice)
            _write_chapter_pdf(source, chapter_slice, chapter_path)
            records.append(
                {
                    **asdict(chapter_slice),
                    "source_start_page": chapter_slice.start_page_index + 1,
                    "source_end_page": chapter_slice.end_page_index + (
                        1 if chapter_slice.end_y > 100 else 0
                    ),
                    "chapter_pdf": str(chapter_path),
                }
            )
    payload = {
        "source_pdf": str(input_path),
        "course_end_page": course_end.page_index,
        "chapter_count": len(records),
        "coverage_status": "complete",
        "chapters": records,
    }
    manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def _data_url(path):
    encoded = base64.b64encode(Path(path).read_bytes()).decode("ascii")
    return f"data:application/pdf;base64,{encoded}"


def _response_format():
    return {
        "type": "json_schema",
        "name": "bold_case_law_chapter",
        "strict": True,
        "schema": BoldChapterExtraction.model_json_schema(),
    }


def _build_input(chapter_record, prompt):
    page_note = (
        f"\nMÉTADONNÉES DE DÉCOUPAGE : la première page locale de ce fichier correspond "
        f"à la page {chapter_record['source_start_page']} du PDF source. "
        f"Le fichier couvre jusqu'à la page {chapter_record['source_end_page']} du PDF source. "
        "Reporte dans source_page le numéro de page du PDF source original.\n"
    )
    return [
        {
            "role": "user",
            "content": [
                {
                    "type": "input_text",
                    "text": prompt + "\n" + BOLD_TECHNICAL_INSTRUCTIONS + page_note,
                },
                {
                    "type": "input_file",
                    "filename": Path(chapter_record["chapter_pdf"]).name,
                    "file_data": _data_url(chapter_record["chapter_pdf"]),
                    "detail": "high",
                },
            ],
        }
    ]


def _parse_extraction(text):
    try:
        return BoldChapterExtraction.model_validate_json(text)
    except ValidationError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if not match:
            raise
        return BoldChapterExtraction.model_validate_json(match.group(0))


async def _extract_one(settings, chapter_record, prompt, retries):
    request_input = _build_input(chapter_record, prompt)
    last_error = None
    for attempt in range(retries):
        try:
            response_text = await _call_openai_responses(
                settings, request_input, text_format=_response_format()
            )
            return _parse_extraction(response_text)
        except Exception as exc:
            last_error = exc
            if attempt + 1 < retries:
                logging.warning(
                    "Échec du chapitre %s (tentative %s/%s) : %s",
                    chapter_record["order"], attempt + 1, retries, exc,
                )
                await asyncio.sleep(2**attempt)
    raise last_error


def _cache_path(cache_dir, order):
    return Path(cache_dir) / f"chapitre_{int(order):03d}.json"


def _write_cache(path, chapter_record, extraction):
    payload = {
        "chapter": chapter_record,
        "result": extraction.model_dump(),
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _sorted_cache_paths(cache_dir):
    return sorted(
        Path(cache_dir).glob("chapitre_*.json"),
        key=lambda path: int(re.search(r"(\d+)$", path.stem).group(1)),
    )


def _read_cache(path):
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    return payload, BoldChapterExtraction.model_validate(payload["result"])


async def _process_chapters(records, settings, cache_dir, prompt, concurrency, retries, force):
    semaphore = asyncio.Semaphore(concurrency)

    async def process(record):
        cache_path = _cache_path(cache_dir, record["order"])
        if cache_path.is_file() and not force:
            return record, "cached", None
        try:
            async with semaphore:
                logging.info("Analyse du chapitre %03d : %s", record["order"], record["title"])
                extraction = await _extract_one(settings, record, prompt, retries)
            _write_cache(cache_path, record, extraction)
            return record, "complete", None
        except Exception as exc:
            return record, "error", str(exc)

    results = []
    tasks = [asyncio.create_task(process(record)) for record in records]
    for done, task in enumerate(asyncio.as_completed(tasks), start=1):
        result = await task
        results.append(result)
        logging.info("Progression %s/%s : chapitre %03d", done, len(tasks), result[0]["order"])
    return results


def _configure_doc_styles(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, (size, color, before, after) in {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
        "Heading 4": (11.5, "1F4D78", 9, 4),
        "Heading 5": (11, "365F7D", 8, 4),
        "Heading 6": (10.5, "365F7D", 7, 3),
    }.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    bullet = document.styles["List Bullet"]
    bullet.base_style = normal
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.25


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def _flatten_rows(cache_dir):
    rows, summaries = [], []
    for cache_path in _sorted_cache_paths(cache_dir):
        payload, extraction = _read_cache(cache_path)
        chapter = payload["chapter"]
        group_count = reference_count = 0
        for entry in extraction.entries:
            for group in entry.groupes:
                rows.append(
                    {
                        "chapter_number": chapter["order"],
                        "chapter_file": Path(chapter["chapter_pdf"]).name,
                        "source_pdf": chapter.get("source_pdf", ""),
                        "partie": entry.partie,
                        "theme": entry.theme,
                        "titre": entry.titre,
                        "chapitre": entry.chapitre,
                        "section": entry.section,
                        "sous_section": entry.sous_section,
                        "references": group.reference_text,
                        "reference_count": len(group.references),
                        "description": group.description,
                        "source_page": group.source_page,
                    }
                )
                group_count += 1
                reference_count += len(group.references)
        summaries.append(
            {
                "chapter_number": chapter["order"],
                "chapter_file": Path(chapter["chapter_pdf"]).name,
                "group_count": group_count,
                "reference_count": reference_count,
                "status": "Traité",
            }
        )
    return rows, summaries


def assemble_bold_case_law_docx(cache_dir, output_docx_path):
    cache_paths = _sorted_cache_paths(cache_dir)
    if not cache_paths:
        raise ValueError("Aucun résultat de chapitre n'a été trouvé.")
    document = Document()
    _configure_doc_styles(document)
    header = document.sections[0].header.paragraphs[0]
    header.text = "Répertoire des jurisprudences en gras"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    _add_page_number(document.sections[0].footer.paragraphs[0])
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Répertoire des jurisprudences signalées en gras")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("0B2545")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    subtitle_run = subtitle.add_run(
        "Références partageant une même portée regroupées sur une seule ligne"
    )
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(10.5)
    subtitle_run.font.color.rgb = RGBColor(90, 90, 90)
    last_headings = [None] * 6
    fields = ["partie", "theme", "titre", "chapitre", "section", "sous_section"]
    group_count = 0
    for cache_path in cache_paths:
        _, extraction = _read_cache(cache_path)
        for entry in extraction.entries:
            values = [getattr(entry, field).strip() for field in fields]
            for index, value in enumerate(values):
                if value and value != last_headings[index]:
                    document.add_heading(value, level=index + 1)
                    last_headings[index] = value
                    for deeper in range(index + 1, 6):
                        last_headings[deeper] = None
            for group in entry.groupes:
                document.add_paragraph(group.formulation, style="List Bullet")
                group_count += 1
    if not group_count:
        document.add_paragraph("Aucune jurisprudence en gras n'a été relevée.")
    output_path = Path(output_docx_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def assemble_bold_case_law_xlsx(cache_dir, output_xlsx_path):
    rows, summaries = _flatten_rows(cache_dir)
    if not summaries:
        raise ValueError("Aucun résultat de chapitre n'a été trouvé.")
    workbook = openpyxl.Workbook()
    results = workbook.active
    results.title = "Jurisprudences"
    summary = workbook.create_sheet("Synthèse")
    headers = [
        "N° chapitre", "Fichier chapitre", "PDF source", "Partie", "Thème", "Titre",
        "Chapitre", "Section", "Sous-section", "Références juridiques groupées",
        "Nombre de références", "Description synthétique", "Page source",
    ]
    keys = [
        "chapter_number", "chapter_file", "source_pdf", "partie", "theme", "titre",
        "chapitre", "section", "sous_section", "references", "reference_count",
        "description", "source_page",
    ]
    results.append(headers)
    for row in rows:
        results.append([row[key] for key in keys])
    summary.append(
        ["N° chapitre", "Fichier chapitre", "Groupes", "Références", "Statut"]
    )
    for row in summaries:
        summary.append(
            [row["chapter_number"], row["chapter_file"], row["group_count"],
             row["reference_count"], row["status"]]
        )
    total_row = summary.max_row + 1
    summary.cell(total_row, 2, "TOTAL")
    summary.cell(total_row, 3, f"=SUM(C2:C{total_row - 1})")
    summary.cell(total_row, 4, f"=SUM(D2:D{total_row - 1})")
    fill = PatternFill("solid", fgColor="1F4D78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10)
    for sheet in (results, summary):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {
        "A": 12, "B": 24, "C": 42, "D": 34, "E": 34, "F": 34, "G": 42,
        "H": 42, "I": 36, "J": 60, "K": 18, "L": 70, "M": 12,
    }.items():
        results.column_dimensions[column].width = width
    for column, width in {"A": 12, "B": 25, "C": 14, "D": 16, "E": 14}.items():
        summary.column_dimensions[column].width = width
    results.row_dimensions[1].height = 30
    summary.row_dimensions[1].height = 30
    if rows:
        table = ExcelTable(displayName="BoldCaseLawResults", ref=f"A1:M{results.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        results.add_table(table)
    summary_table = ExcelTable(
        displayName="BoldChapterSummary", ref=f"A1:E{total_row - 1}"
    )
    summary_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    summary.add_table(summary_table)
    for column in (2, 3, 4):
        summary.cell(total_row, column).font = Font(name="Calibri", bold=True)
    output_path = Path(output_xlsx_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def assemble_bold_case_law_outputs(cache_dir, output_docx_path, output_xlsx_path=None):
    docx_path = assemble_bold_case_law_docx(cache_dir, output_docx_path)
    xlsx_path = assemble_bold_case_law_xlsx(
        cache_dir, output_xlsx_path or Path(output_docx_path).with_suffix(".xlsx")
    )
    return docx_path, xlsx_path


def extract_bold_case_law_pdf(
    input_pdf_path,
    output_docx_path,
    output_xlsx_path=None,
    chapters_dir=None,
    manifest_path=None,
    cache_dir=None,
    open_ai_key=None,
    model=None,
    concurrency=2,
    retries=3,
    env_path=".env",
    prompt=BOLD_CASE_LAW_PROMPT,
    force=False,
    from_chapter=None,
    to_chapter=None,
    reuse_chapters=False,
):
    input_path = Path(input_pdf_path).resolve()
    output_path = Path(output_docx_path).resolve()
    chapters_path = (
        Path(chapters_dir).resolve()
        if chapters_dir
        else output_path.parent / f"{output_path.stem}_chapters"
    )
    manifest = (
        Path(manifest_path).resolve()
        if manifest_path
        else chapters_path / "chapters_manifest.json"
    )
    if reuse_chapters:
        if not manifest.exists():
            raise FileNotFoundError(
                f"Manifeste de chapitres introuvable pour la reprise : {manifest}"
            )
        existing_payload = json.loads(manifest.read_text(encoding="utf-8"))
        missing = [
            record.get("chapter_pdf", "")
            for record in existing_payload.get("chapters", [])
            if not Path(record.get("chapter_pdf", "")).exists()
        ]
        if missing:
            raise FileNotFoundError(
                f"{len(missing)} PDF de chapitre manquent ; relancez sans --reuse-chapters."
            )
    else:
        manifest = split_pdf_by_detected_chapters(input_path, chapters_path, manifest)
    payload = json.loads(manifest.read_text(encoding="utf-8"))
    records = payload["chapters"]
    for record in records:
        record["source_pdf"] = str(input_path)
    if from_chapter is not None:
        records = [record for record in records if record["order"] >= from_chapter]
    if to_chapter is not None:
        records = [record for record in records if record["order"] <= to_chapter]
    if not records:
        raise ValueError("Aucun chapitre ne correspond à la sélection.")
    cache_path = (
        Path(cache_dir).resolve()
        if cache_dir
        else output_path.parent / f"{output_path.stem}_results"
    )
    settings = _resolve_openai_settings(open_ai_key=open_ai_key, model=model, env_path=env_path)
    if not settings["api_key"]:
        raise ValueError("La clé OpenAI est absente.")
    results = asyncio.run(
        _process_chapters(
            records, settings, cache_path, prompt, concurrency, retries, force
        )
    )
    errors = [
        (record["order"], error)
        for record, status, error in results
        if status == "error"
    ]
    if errors:
        details = " | ".join(f"chapitre {order}: {error}" for order, error in errors)
        raise RuntimeError(
            f"{len(errors)} chapitre(s) en erreur. Les réussites sont conservées : {details}"
        )
    return assemble_bold_case_law_outputs(cache_path, output_path, output_xlsx_path)
