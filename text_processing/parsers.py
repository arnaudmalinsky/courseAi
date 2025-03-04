import io
from pathlib import Path
import logging

from docx import Document
from openpyxl import Workbook
from langchain.text_splitter import RecursiveCharacterTextSplitter

class CorpusParser:
    def __init__(self,folder_path, output_path, max_length):
        self.folder_path = Path(folder_path)
        self.get_file_path()
        self.output_path=output_path
        self.max_length=max_length

    def get_file_path(self):
        # self.file_path_list = list(self.folder_path.glob('**/*.docx'))
        self.file_path_list = list(self.folder_path.glob('*.docx'))

    def parse_all_documents(self):
        wb = Workbook()
        worksheet = wb.active
        worksheet.title = "Document Structure"
        worksheet.append(["Folder","Filename","Index", "Type", "Text", "Character Index", "Title Context","Title lvl2","Title lvl3"])
        for file_path in self.file_path_list:
            new_document = (
                DocumentParser(file_path, worksheet,self.max_length).parse_document()
            )
        wb.save(self.output_path)


class DocumentParser:
    """
    Could be used to store 'structure' from docx package
    """

    def __init__(self, filename, worksheet,max_length):
        self.filename = filename
        self.worksheet=worksheet
        self.max_lenght=max_length
        
        
        
    def parse_document(self):
        char_index = 0  # Running character index for the document
        heading_context = {1: "", 2: "", 3: ""}  # Track context for heading levels
        
        with open(self.filename, 'rb') as f:
            source_stream = io.BytesIO(f.read())
            document = Document(source_stream)
            for i, paragraph in enumerate(document.paragraphs):
                new_paragraph = ParagraphElement(self.filename, i, paragraph, char_index,heading_context, self.max_lenght, self.worksheet)
                heading_context = new_paragraph.update_context()
                new_paragraph.parse_and_save()
                char_index=new_paragraph.get_char_index()
                


class ParagraphElement:

    """
    Could be used to store element from docx package
    """

    def __init__(
            self, 
            filename, 
            index,
            paragraph,
            char_index,
            heading_context,
            max_length,
            worksheet
        ):
        self.worksheet=worksheet
        self.max_lenght=max_length
        self.heading_level_1_context=""
        self.heading_level_2_context=""
        self.heading_level_3_context=""
        self.filename = filename
        self.index=index
        self.paragraph = paragraph
        self.text = self.paragraph.text.strip()
        self.character_index = char_index
        self.heading_context=heading_context
        self.style_name = self.paragraph.style.name
        self.get_text_type()

        
    def get_text_type(self):
        if "Heading" in self.style_name:
            self.text_type = "title"
        else:
            self.text_type = "paragraph"
    
    def update_context(self):
        if self.text_type == 'title':
            try:
                heading_level = int(self.style_name.split()[-1])  # Extract numeric level
                if heading_level in self.heading_context.keys():
                    self.heading_context[heading_level] = self.text
                    # Clear lower-level contexts
                    for lvl in range(heading_level + 1, 4):
                        self.heading_context[lvl] = ""
            except ValueError:
                heading_level = None
        return self.heading_context

    def update_cursor(self):
        self.character_index += len(self.text) + 1

    def get_char_index(self):
        return self.character_index 
    
    @staticmethod
    def _breakdown_paragraph(text, max_length=600, overlap=0):
        # paragraphs = []
        # while len(text) > max_length:
            # split_point = max_length - overlap
            # paragraphs.append(text[:split_point].strip())
            # text = text[split_point:].strip()
        # paragraphs.append(text)
        logging.warning(f"input is too long compared to max_length: {max_length}")

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=max_length,  # Max chunk length
            chunk_overlap=overlap,  # Overlap between chunks to preserve context
            separators=["\n\n", "\n", ". ", " "],  # Order of splitting
        )

        chunks = splitter.split_text(text)
            
        return chunks

    def parse_and_save(self):
        sub_paragraphs = self._breakdown_paragraph(self.text, max_length=self.max_lenght)
        for sub_paragraph in sub_paragraphs:
            self.text=sub_paragraph
            self.character_index=self.character_index
            self.heading_level_1_context=self.heading_context[1]
            self.heading_level_2_context=self.heading_context[2]
            self.heading_level_3_context=self.heading_context[3]
            self.update_cursor() 
            self.save()
        
    def __repr__(self):
        print(f"Docx file: {self.filename}")
        print(f"Type: {self.text_type.capitalize()}, Index: {self.index}")
        print(f"Text: {self.text}")
        print(f"Character Index: {self.character_index}")
        print(f"Heading Level 1 Context: {self.heading_level_1_context}")
        print(f"Heading Level 2 Context: {self.heading_level_2_context}")
        print(f"Heading Level 3 Context: {self.heading_level_3_context}")
        print("---")

    def save(self):
        # Write structure data
        self.worksheet.append([
            str(self.filename.parents[0]),
            self.filename.stem,
            self.index,
            self.text_type,
            self.text,
            self.character_index,
            self.heading_level_1_context,
            self.heading_level_2_context,
            self.heading_level_3_context,
        ])
