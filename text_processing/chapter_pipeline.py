import asyncio
import hashlib
import logging
import re
import time
from copy import deepcopy
from pathlib import Path

import openpyxl
import httpx
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from prompting.sumup_prompt import CHAPTER_SUMMARY_PROMPT


MANIFEST_HEADERS = [
    "chapter_id", "order", "section_type", "partie", "theme", "title", "source_docx",
    "source_length", "source_block_count", "source_heading_count",
    "completeness_status", "status", "summary", "summary_length",
    "reduction_rate", "validation", "error",
]

CHAPTER_RE = re.compile(
    r"^\s*CHAPITRE\s+(?:[IVXLCDM]+|\d+(?:ER|E|ÈME)?|PREMIER|PREMI[ÈE]RE|"
    r"UNIQUE|PR[ÉE]LIMINAIRE)\b",
    re.IGNORECASE,
)
PARTIE_RE = re.compile(
    r"^\s*PARTIE\s+(?:[IVXLCDM]+|\d+(?:ER|E|ÈME)?|PREMI[ÈE]RE)\b",
    re.IGNORECASE,
)
THEME_RE = re.compile(
    r"^\s*TH[ÈE]ME\s+(?:[IVXLCDM]+|\d+(?:ER|E|ÈME)?|PREMIER)\b",
    re.IGNORECASE,
)
MARKDOWN_HEADING_RE = re.compile(r"^(#{1,9})\s+(.+?)\s*$")
BULLET_RE = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)")
TOC_ENTRY_RE = re.compile(r"\.{5,}\s*\d+\s*$")
TOC_PAGE_ONLY_RE = re.compile(r"^\s*Chapitre\b.*\s+\d{1,3}\s*$")


def _load_env_file(env_path=".env"):
    values = {}
    path = Path(env_path)
    if not path.is_file():
        return values
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip("\"'")
    return values


def _resolve_openai_settings(
    open_ai_key=None,
    model=None,
    reasoning_effort=None,
    verbosity=None,
    reasoning_summary=None,
    env_path=".env",
):
    env_values = _load_env_file(env_path)
    key = (
        open_ai_key
        or env_values.get("OPENAI_API_KEY")
        or env_values.get("OPENAI_KEY")
    )
    resolved_reasoning = (
        reasoning_effort
        or env_values.get("MODEL_REASONING")
        or env_values.get("MODEL_REASONNING")
    )
    if resolved_reasoning and resolved_reasoning.lower() == "standard":
        logging.warning(
            "MODEL_REASONNING=standard is mapped to the supported API value 'medium'."
        )
        resolved_reasoning = "medium"
    return {
        "api_key": key,
        "model": model or env_values.get("OPENAI_MODEL") or "gpt-4o",
        "reasoning_effort": resolved_reasoning,
        "verbosity": verbosity or env_values.get("MODEL_VERBOSITY"),
        "reasoning_summary": reasoning_summary or env_values.get("MODEL_SUMMARY"),
    }


def _iter_block_items(document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _block_text(block):
    if isinstance(block, Paragraph):
        return block.text
    return "\n".join(
        "\t".join(cell.text for cell in row.cells)
        for row in block.rows
    )


def _normalized_text(text):
    return " ".join((text or "").split())


def _block_signature(block):
    block_type = "paragraph" if isinstance(block, Paragraph) else "table"
    return f"{block_type}:{_normalized_text(_block_text(block))}"


def _sequence_hash(signatures):
    return hashlib.sha256("\n".join(signatures).encode("utf-8")).hexdigest()


def _paragraph_heading_level(paragraph):
    if not isinstance(paragraph, Paragraph):
        return None
    match = re.search(r"(?:Heading|Titre)\s+(\d+)$", paragraph.style.name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _is_chapter_start(block):
    return isinstance(block, Paragraph) and bool(CHAPTER_RE.match(block.text.strip()))


def _upper_heading_kind(block):
    if not isinstance(block, Paragraph):
        return None
    text = block.text.strip()
    if PARTIE_RE.match(text):
        return "partie"
    if THEME_RE.match(text):
        return "theme"
    return None


def _chapter_contexts(blocks, chapter_starts):
    contexts = {}
    partie = ""
    theme = ""
    chapter_start_set = set(chapter_starts)
    for index, block in enumerate(blocks):
        kind = _upper_heading_kind(block)
        if kind == "partie":
            partie = block.text.strip()
            theme = ""
        elif kind == "theme":
            theme = block.text.strip()
        if index in chapter_start_set:
            contexts[index] = {"partie": partie, "theme": theme}
    return contexts


def _chapter_boundaries(blocks, chapter_starts):
    boundaries = []
    first_chapter = chapter_starts[0]
    upper_before_first = [
        index
        for index in range(first_chapter)
        if _upper_heading_kind(blocks[index])
    ]
    if upper_before_first:
        last_partie = next(
            (
                index
                for index in reversed(upper_before_first)
                if _upper_heading_kind(blocks[index]) == "partie"
            ),
            None,
        )
        boundaries.append(last_partie if last_partie is not None else upper_before_first[-1])
    else:
        boundaries.append(first_chapter)

    for previous_chapter, chapter_start in zip(chapter_starts, chapter_starts[1:]):
        upper_between = [
            index
            for index in range(previous_chapter + 1, chapter_start)
            if _upper_heading_kind(blocks[index])
        ]
        boundaries.append(upper_between[0] if upper_between else chapter_start)
    return boundaries


def _safe_filename(order, title):
    slug = re.sub(r"[^\w.-]+", "_", title, flags=re.UNICODE).strip("_.")
    return f"{order:03d}_{slug[:90] or 'chapitre'}.docx"


def _is_table_of_contents_entry(title):
    text = str(title or "").strip()
    return bool(TOC_ENTRY_RE.search(text) or TOC_PAGE_ONLY_RE.match(text))


def _copy_blocks_to_docx(blocks, output_path):
    target = Document()
    target_body = target.element.body
    target_sect_pr = target_body.sectPr
    for block in blocks:
        target_body.insert(target_body.index(target_sect_pr), deepcopy(block._element))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    target.save(output_path)


def _write_manifest(rows, manifest_path):
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Chapters"
    worksheet.append(MANIFEST_HEADERS)
    for row in rows:
        worksheet.append([row.get(header, "") for header in MANIFEST_HEADERS])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    widths = {
        "A": 14, "B": 9, "C": 45, "D": 45, "E": 45, "F": 60, "G": 16,
        "H": 20, "I": 20, "J": 22, "K": 16, "L": 80, "M": 16, "N": 16,
        "O": 45, "P": 45,
    }
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(manifest_path)


def split_docx_by_chapter(source_docx_path, output_dir, manifest_path=None):
    source_path = Path(source_docx_path).resolve()
    output_path = Path(output_dir).resolve()
    manifest_path = (
        Path(manifest_path).resolve()
        if manifest_path
        else output_path / "chapters_manifest.xlsx"
    )

    document = Document(source_path)
    blocks = list(_iter_block_items(document))
    chapter_starts = [index for index, block in enumerate(blocks) if _is_chapter_start(block)]
    if not chapter_starts:
        raise ValueError(
            "Aucun titre commençant par « CHAPITRE » n'a été détecté dans le DOCX."
        )

    contexts = _chapter_contexts(blocks, chapter_starts)
    boundaries = _chapter_boundaries(blocks, chapter_starts)
    ranges = []
    if boundaries[0] > 0:
        ranges.append(("preamble", 0, boundaries[0], None, "Préambule"))
    for chapter_index, chapter_start in enumerate(chapter_starts):
        start = boundaries[chapter_index]
        end = boundaries[chapter_index + 1] if chapter_index + 1 < len(boundaries) else len(blocks)
        ranges.append(
            ("chapter", start, end, chapter_start, blocks[chapter_start].text.strip())
        )

    rows = []
    exported_paths = []
    for order, (kind, start, end, chapter_start, title) in enumerate(ranges):
        selected_blocks = blocks[start:end]
        filename = "000_preambule.docx" if kind == "preamble" else _safe_filename(order, title)
        chapter_path = output_path / filename
        _copy_blocks_to_docx(selected_blocks, chapter_path)
        exported_paths.append(chapter_path)
        section_type = (
            "preamble"
            if kind == "preamble"
            else "table_of_contents"
            if _is_table_of_contents_entry(title)
            else "chapter"
        )
        rows.append({
            "chapter_id": "preamble" if kind == "preamble" else f"chapter_{order:03d}",
            "order": order,
            "section_type": section_type,
            "partie": contexts[chapter_start]["partie"] if chapter_start is not None else "",
            "theme": contexts[chapter_start]["theme"] if chapter_start is not None else "",
            "title": title,
            "source_docx": str(chapter_path),
            "source_length": sum(len(_block_text(block)) for block in selected_blocks),
            "source_block_count": len(selected_blocks),
            "source_heading_count": sum(
                _paragraph_heading_level(block) is not None for block in selected_blocks
            ),
            "completeness_status": "pending",
            "status": "excluded" if section_type == "table_of_contents" else "pending",
        })

    source_signatures = [_block_signature(block) for block in blocks]
    exported_signatures = []
    for exported_path in exported_paths:
        exported_document = Document(exported_path)
        exported_signatures.extend(
            _block_signature(block) for block in _iter_block_items(exported_document)
        )
    complete = (
        len(source_signatures) == len(exported_signatures)
        and source_signatures == exported_signatures
        and _sequence_hash(source_signatures) == _sequence_hash(exported_signatures)
    )
    for row in rows:
        row["completeness_status"] = "complete" if complete else "invalid"
        if not complete:
            row["error"] = (
                "AVERTISSEMENT : la concaténation des exports ne reproduit pas "
                "exactement le DOCX source."
            )

    _write_manifest(rows, manifest_path)
    if not complete:
        logging.warning(
            "Completeness validation failed; processing remains available. Manifest: %s",
            manifest_path,
        )
    else:
        logging.warning("Completeness check passed; manifest saved to %s", manifest_path)
    return manifest_path


def _docx_to_markdown(path):
    document = Document(path)
    lines, headings = [], []
    source_has_bullets = False
    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            if _upper_heading_kind(block):
                continue
            level = _paragraph_heading_level(block)
            if level:
                headings.append((level, text))
                lines.append(f"{'#' * level} {text}")
            else:
                is_list = "list" in block.style.name.lower() or bool(BULLET_RE.match(text))
                source_has_bullets = source_has_bullets or is_list
                lines.append(text)
        else:
            lines.append(_block_text(block))
    return "\n\n".join(lines), headings, source_has_bullets


def _extract_markdown_headings(text):
    result = []
    for line in (text or "").splitlines():
        match = MARKDOWN_HEADING_RE.match(line.strip())
        if match:
            result.append((len(match.group(1)), match.group(2).strip()))
    return result


def validate_summary(
    source_text,
    source_headings,
    source_has_bullets,
    summary,
    enforce_reduction=True,
):
    errors = []
    if _extract_markdown_headings(summary) != source_headings:
        errors.append("Les titres, leur ordre ou leur niveau ne correspondent pas à la source.")
    source_length = len(_normalized_text(source_text))
    summary_length = len(_normalized_text(summary))
    reduction_rate = 1 - (summary_length / source_length) if source_length else 0
    if enforce_reduction and source_length and reduction_rate < 0.5:
        errors.append("La réduction est inférieure à 50 %.")
    if not source_has_bullets and any(BULLET_RE.match(line) for line in summary.splitlines()):
        errors.append("Le résumé ajoute des listes absentes de la source.")
    forbidden = ("Idée principale", "Jurisprudence :", "Portée :", "Norme juridique")
    if any(label.lower() in summary.lower() for label in forbidden):
        errors.append("Le résumé contient une catégorie interdite.")
    return errors, source_length, summary_length, reduction_rate


def _load_manifest(manifest_path):
    workbook = openpyxl.load_workbook(manifest_path)
    worksheet = workbook["Chapters"]
    headers = [cell.value for cell in worksheet[1]]
    for header in MANIFEST_HEADERS:
        if header not in headers:
            headers.append(header)
            worksheet.cell(row=1, column=len(headers), value=header)
    rows = [
        (row_number, dict(zip(headers, values)))
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        )
    ]
    return workbook, worksheet, headers, rows


def _extract_response_text(response_data):
    if response_data.get("output_text"):
        return response_data["output_text"]
    texts = []
    for item in response_data.get("output", []):
        if item.get("type") != "message":
            continue
        for content in item.get("content", []):
            if content.get("type") == "output_text" and content.get("text"):
                texts.append(content["text"])
    if not texts:
        raise ValueError("La réponse OpenAI ne contient aucun texte exploitable.")
    return "\n".join(texts)


async def _call_openai_responses(settings, prompt, text_format=None):
    payload = {
        "model": settings["model"],
        "input": prompt,
    }
    reasoning = {}
    if settings.get("reasoning_effort"):
        reasoning["effort"] = settings["reasoning_effort"]
    if settings.get("reasoning_summary"):
        reasoning["summary"] = settings["reasoning_summary"]
    if reasoning:
        payload["reasoning"] = reasoning
    text_options = {}
    if settings.get("verbosity"):
        text_options["verbosity"] = settings["verbosity"]
    if text_format:
        text_options["format"] = text_format
    if text_options:
        payload["text"] = text_options

    timeout = httpx.Timeout(600.0, connect=30.0)
    async with httpx.AsyncClient(timeout=timeout) as client:
        response = await client.post(
            "https://api.openai.com/v1/responses",
            headers={
                "Authorization": f"Bearer {settings['api_key']}",
                "Content-Type": "application/json",
            },
            json=payload,
        )
        response.raise_for_status()
        return _extract_response_text(response.json())


async def _summarize_one(settings, semaphore, source_text, retries, correction=None):
    technical_instruction = """
INSTRUCTIONS TECHNIQUES DE RESTITUTION :
- Le contenu transmis représente une section autonome du document.
- Les titres sont signalés par des marqueurs Markdown #, ##, ###, etc.
- Reproduis chaque titre avec exactement le même marqueur, le même texte et dans le même ordre.
- Restitue uniquement la fiche, sans introduction ni commentaire sur ton travail.
"""
    correction_text = (
        f"\nCORRECTIONS OBLIGATOIRES APRÈS CONTRÔLE :\n{correction}\n"
        if correction
        else ""
    )
    prompt = (
        f"{CHAPTER_SUMMARY_PROMPT}\n{technical_instruction}{correction_text}"
        f"\nDOCUMENT À TRAITER :\n{source_text}"
    )
    last_error = None
    for attempt in range(retries):
        try:
            async with semaphore:
                return await _call_openai_responses(settings, prompt)
        except Exception as exc:
            last_error = exc
            if attempt + 1 < retries:
                logging.warning(
                    "Échec de l'appel OpenAI (tentative %s/%s) : %s. Nouvelle tentative...",
                    attempt + 1,
                    retries,
                    exc,
                )
                await asyncio.sleep(2 ** attempt)
    raise last_error


async def _summarize_pending_rows(
    rows, settings, concurrency, retries, on_result=None
):
    semaphore = asyncio.Semaphore(concurrency)

    async def process(row_number, row):
        logging.info(
            "Démarrage [%s | ordre %s] %s",
            row.get("chapter_id"),
            row.get("order"),
            row.get("title"),
        )
        source_text, headings, source_has_bullets = _docx_to_markdown(row["source_docx"])
        try:
            summary = await _summarize_one(
                settings, semaphore, source_text, retries
            )
            validation, source_length, summary_length, reduction_rate = validate_summary(
                source_text,
                headings,
                source_has_bullets,
                summary,
                enforce_reduction=False,
            )
            return row_number, {
                "status": "complete",
                "summary": summary,
                "source_length": source_length,
                "summary_length": summary_length,
                "reduction_rate": reduction_rate,
                "validation": "OK" if not validation else "WARNING: " + " | ".join(validation),
                "error": "",
            }, None
        except Exception as exc:
            return row_number, None, str(exc)

    completed_results = []
    tasks = [asyncio.create_task(process(row_number, row)) for row_number, row in rows]
    for task in asyncio.as_completed(tasks):
        result = await task
        completed_results.append(result)
        if on_result:
            on_result(result)
    return completed_results


def summarize_chapter_manifest(
    manifest_path,
    open_ai_key=None,
    model=None,
    concurrency=4,
    retries=3,
    retry_invalid=False,
    reasoning_effort=None,
    verbosity=None,
    reasoning_summary=None,
    env_path=".env",
    chapter_ids=None,
    orders=None,
    from_order=None,
    to_order=None,
    force=False,
):
    settings = _resolve_openai_settings(
        open_ai_key,
        model,
        reasoning_effort,
        verbosity,
        reasoning_summary,
        env_path,
    )
    if not settings["api_key"]:
        raise ValueError("La clé OpenAI est absente.")
    if concurrency < 1:
        raise ValueError("concurrency doit être supérieur ou égal à 1.")
    manifest_path = Path(manifest_path).resolve()
    workbook, worksheet, headers, rows = _load_manifest(manifest_path)
    header_columns = {header: index + 1 for index, header in enumerate(headers)}
    requested_ids = {str(value) for value in (chapter_ids or [])}
    requested_orders = {int(value) for value in (orders or [])}
    if from_order is not None:
        from_order = int(from_order)
    if to_order is not None:
        to_order = int(to_order)
    if (
        from_order is not None
        and to_order is not None
        and from_order > to_order
    ):
        raise ValueError("--from-order doit être inférieur ou égal à --to-order.")
    if requested_ids:
        available_ids = {str(row.get("chapter_id")) for _, row in rows}
        missing_ids = requested_ids - available_ids
        if missing_ids:
            raise ValueError(
                "chapter_id introuvable dans le manifeste : "
                + ", ".join(sorted(missing_ids))
            )
    if requested_orders:
        available_orders = {int(row.get("order")) for _, row in rows if row.get("order") is not None}
        missing_orders = requested_orders - available_orders
        if missing_orders:
            raise ValueError(
                "Ordre introuvable dans le manifeste : "
                + ", ".join(str(value) for value in sorted(missing_orders))
            )

    def is_selected(row):
        row_order = int(row.get("order"))
        in_range = (
            (from_order is None or row_order >= from_order)
            and (to_order is None or row_order <= to_order)
        )
        has_range = from_order is not None or to_order is not None
        if requested_ids or requested_orders or has_range:
            return (
                str(row.get("chapter_id")) in requested_ids
                or row_order in requested_orders
                or (has_range and in_range)
            )
        return True

    pending = [
        (row_number, row)
        for row_number, row in rows
        if is_selected(row)
        and row.get("status") != "excluded"
        and (
            force
            or row.get("status") in (None, "", "pending", "processing", "error", "invalid")
            or (retry_invalid and row.get("status") == "invalid")
        )
    ]
    logging.info(
        "Manifeste chargé : %s ligne(s), %s chapitre(s) sélectionné(s) pour traitement.",
        len(rows),
        len(pending),
    )
    if not pending:
        logging.info(
            "Aucun chapitre à traiter. Utilisez --force pour relancer des chapitres terminés."
        )
        return manifest_path

    for row_number, _ in pending:
        worksheet.cell(row=row_number, column=header_columns["status"], value="processing")
    workbook.save(manifest_path)
    logging.info(
        "Statut 'processing' enregistré pour %s chapitre(s) dans %s.",
        len(pending),
        manifest_path,
    )

    started_at = time.monotonic()
    completed_count = 0
    success_count = 0
    error_count = 0
    row_lookup = {row_number: row for row_number, row in pending}

    def save_result(result):
        nonlocal completed_count, success_count, error_count
        row_number, values, error = result
        if error:
            values = {"status": "error", "error": error}
        for field, value in values.items():
            worksheet.cell(row=row_number, column=header_columns[field], value=value)
        workbook.save(manifest_path)
        completed_count += 1
        if error:
            error_count += 1
        else:
            success_count += 1
        elapsed = time.monotonic() - started_at
        average = elapsed / completed_count
        remaining_seconds = average * (len(pending) - completed_count)
        progress = completed_count / len(pending) * 100
        row = row_lookup[row_number]
        validation = values.get("validation")
        logging.info(
            "Progression %s/%s (%.1f%%) | %s | [%s | ordre %s] %s | "
            "écoulé %.1f min | restant estimé %.1f min%s",
            completed_count,
            len(pending),
            progress,
            values.get("status"),
            row.get("chapter_id"),
            row.get("order"),
            row.get("title"),
            elapsed / 60,
            remaining_seconds / 60,
            f" | {validation}" if validation and validation != "OK" else "",
        )

    asyncio.run(
        _summarize_pending_rows(
            pending,
            settings,
            concurrency,
            retries,
            on_result=save_result,
        )
    )
    elapsed = time.monotonic() - started_at
    logging.info(
        "Traitement terminé : %s succès, %s erreur(s), %s chapitre(s), durée %.1f min. "
        "Résultats sauvegardés dans %s.",
        success_count,
        error_count,
        len(pending),
        elapsed / 60,
        manifest_path,
    )
    return manifest_path


def _add_markdown_runs(paragraph, text):
    for part in re.split(r"(\*\*.*?\*\*|\*.*?\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def assemble_chapter_summaries(manifest_path, output_docx_path, title=None):
    _, _, _, rows = _load_manifest(Path(manifest_path).resolve())
    rows_with_summary = [
        (row_number, row)
        for row_number, row in rows
        if str(row.get("summary") or "").strip()
    ]
    if not rows_with_summary:
        raise ValueError("Assemblage impossible : aucun résumé n'est disponible.")
    total_source_length = sum(int(row.get("source_length") or 0) for _, row in rows_with_summary)
    total_summary_length = sum(int(row.get("summary_length") or 0) for _, row in rows_with_summary)
    global_reduction = (
        1 - (total_summary_length / total_source_length)
        if total_source_length
        else 0
    )
    if total_source_length and global_reduction < 0.5:
        logging.warning(
            "La réduction globale est inférieure à 50 % (validation non bloquante)."
        )
    document = Document()
    if title:
        document.add_heading(title, level=0)
    previous_partie = None
    previous_theme = None
    for _, row in sorted(rows_with_summary, key=lambda item: item[1]["order"]):
        partie = str(row.get("partie") or "").strip()
        theme = str(row.get("theme") or "").strip()
        if partie and partie != previous_partie:
            document.add_heading(partie, level=1)
            previous_partie = partie
            previous_theme = None
        if theme and theme != previous_theme:
            document.add_heading(theme, level=2)
            previous_theme = theme
        for raw_line in (row.get("summary") or "").splitlines():
            line = raw_line.strip()
            if not line:
                continue
            heading_match = MARKDOWN_HEADING_RE.match(line)
            if heading_match:
                heading_text = heading_match.group(2).strip()
                if heading_text in (partie, theme):
                    continue
                document.add_heading(
                    heading_text, level=min(len(heading_match.group(1)), 9)
                )
            else:
                _add_markdown_runs(document.add_paragraph(), line)
    output_path = Path(output_docx_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
