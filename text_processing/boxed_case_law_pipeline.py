import asyncio
import base64
import json
import logging
import re
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from .chapter_pipeline import _call_openai_responses, _resolve_openai_settings


DEFAULT_PROMPT = """À partir de ce manuel PDF, dresse un répertoire exhaustif de toutes les jurisprudences figurant dans les encadrés du manuel, dans la photo jointe tu trouveras un exemple, il contient deux références juridiques : arrêt Tête du 28 juillet 2000, CE, avis M. Provin du 27 mai 2005
Consignes impératives :
Ne retiens que les jurisprudences présentes dans un encadré.
N'ajoute jamais une jurisprudence qui apparaît uniquement dans le texte courant.
N'ajoute aucune jurisprudence issue de tes connaissances personnelles.
Respecte strictement le plan du manuel.
Reprends exactement les intitulés des Parties, Thèmes, Titres, Chapitres, Sections et Sous-sections.
Classe chaque décision dans la partie du plan où elle apparaît.
Pour chaque jurisprudence, adopte exclusivement la présentation suivante :
CE, Ass., 17 février 1950, Dame Lamotte : consacre comme principe général du droit le recours pour excès de pouvoir ouvert même sans texte contre toute décision administrative."""

TECHNICAL_INSTRUCTIONS = """

INSTRUCTIONS TECHNIQUES DE SORTIE :
- Analyse visuellement toutes les pages du PDF, notamment les bordures et fonds qui délimitent les encadrés.
- La photo jointe est uniquement un exemple visuel de ce qui constitue un encadré ; n'ajoute pas ses décisions sauf si elles figurent aussi dans le PDF traité.
- Une décision n'est admissible que si sa référence apparaît à l'intérieur des limites visuelles d'un encadré.
- Utilise le texte entourant chaque référence dans le même encadré pour synthétiser sa portée après les deux-points.
- N'invente aucun niveau de plan absent du PDF. Laisse la chaîne vide si un niveau n'est pas visible.
- Reproduis les intitulés du plan caractère pour caractère, sans les corriger ni les reformuler.
- Une entrée peut regrouper plusieurs décisions appartenant exactement au même emplacement du plan.
- Le champ « reference » contient uniquement la référence : juridiction, formation éventuelle, date et nom de la décision, sans deux-points final.
- Le champ « description » contient uniquement la portée synthétique tirée du contexte de l'encadré, sans répéter la référence.
- La concaténation « reference : description » doit respecter exactement la présentation imposée.
- N'ajoute aucun commentaire, avertissement ou texte hors du format structuré demandé.
"""


class CaseLawItem(BaseModel):
    model_config = ConfigDict(extra="forbid")

    reference: str = Field(
        description="Juridiction, formation éventuelle, date et nom de la décision"
    )
    description: str = Field(description="Portée synthétique de la décision")
    source_page: int | None = Field(description="Page du PDF du chapitre")

    @property
    def formulation(self):
        return f"{self.reference.rstrip(' :')} : {self.description.lstrip(' :')}"


class PlanEntry(BaseModel):
    model_config = ConfigDict(extra="forbid")

    partie: str
    theme: str
    titre: str
    chapitre: str
    section: str
    sous_section: str
    jurisprudences: list[CaseLawItem]


class ChapterExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")

    entries: list[PlanEntry]


def _data_url(path, mime_type):
    encoded = base64.b64encode(Path(path).read_bytes()).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def _chapter_number(path):
    match = re.search(r"(\d+)$", path.stem)
    return int(match.group(1)) if match else 10**9


def _response_format():
    schema = ChapterExtraction.model_json_schema()
    return {
        "type": "json_schema",
        "name": "boxed_case_law_chapter",
        "strict": True,
        "schema": schema,
    }


def _build_input(pdf_path, example_image_path, prompt):
    return [
        {
            "role": "user",
            "content": [
                {"type": "input_text", "text": prompt + TECHNICAL_INSTRUCTIONS},
                {
                    "type": "input_image",
                    "image_url": _data_url(example_image_path, "image/jpeg"),
                    "detail": "high",
                },
                {
                    "type": "input_file",
                    "filename": Path(pdf_path).name,
                    "file_data": _data_url(pdf_path, "application/pdf"),
                    "detail": "high",
                },
            ],
        }
    ]


def _parse_extraction(text):
    try:
        return ChapterExtraction.model_validate_json(text)
    except ValidationError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if not match:
            raise
        return ChapterExtraction.model_validate_json(match.group(0))


def _normalize_cached_result(result):
    """Convertit les anciens caches `formulation` vers le nouveau schéma."""
    normalized = json.loads(json.dumps(result, ensure_ascii=False))
    for entry in normalized.get("entries", []):
        for case in entry.get("jurisprudences", []):
            if "reference" in case and "description" in case:
                continue
            formulation = str(case.pop("formulation", "")).strip()
            reference, separator, description = formulation.partition(" : ")
            if not separator:
                reference, separator, description = formulation.partition(":")
            case["reference"] = reference.strip()
            case["description"] = description.strip() if separator else ""
    return normalized


async def _extract_one(settings, pdf_path, example_image_path, prompt, retries):
    last_error = None
    request_input = _build_input(pdf_path, example_image_path, prompt)
    for attempt in range(retries):
        try:
            response_text = await _call_openai_responses(
                settings,
                request_input,
                text_format=_response_format(),
            )
            return _parse_extraction(response_text)
        except Exception as exc:
            last_error = exc
            if attempt + 1 < retries:
                logging.warning(
                    "Échec de %s (tentative %s/%s) : %s",
                    Path(pdf_path).name,
                    attempt + 1,
                    retries,
                    exc,
                )
                await asyncio.sleep(2**attempt)
    raise last_error


def _cache_path(cache_dir, pdf_path):
    return Path(cache_dir) / f"{Path(pdf_path).stem}.json"


def _write_cache(path, pdf_path, extraction):
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source_pdf": str(Path(pdf_path).resolve()),
        "result": extraction.model_dump(),
    }
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def _read_cache(path):
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    return ChapterExtraction.model_validate(_normalize_cached_result(payload["result"]))


def _configure_doc_styles(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
        "Heading 4": (11.5, "1F4D78", 9, 4),
        "Heading 5": (11, "365F7D", 8, 4),
        "Heading 6": (10.5, "365F7D", 7, 3),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    style = document.styles["List Bullet"]
    style.base_style = normal
    style.paragraph_format.left_indent = Inches(0.375)
    style.paragraph_format.first_line_indent = Inches(-0.188)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])


def _sorted_cache_paths(cache_dir):
    return sorted(Path(cache_dir).glob("chapitre_*.json"), key=_chapter_number)


def _flatten_case_law_rows(cache_dir):
    rows = []
    chapter_counts = []
    for cache_path in _sorted_cache_paths(cache_dir):
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
        extraction = ChapterExtraction.model_validate(
            _normalize_cached_result(payload["result"])
        )
        chapter_number = _chapter_number(cache_path)
        count = 0
        for entry in extraction.entries:
            for case in entry.jurisprudences:
                rows.append(
                    {
                        "chapter_number": chapter_number,
                        "chapter_file": cache_path.stem + ".pdf",
                        "source_pdf": payload.get("source_pdf", ""),
                        "partie": entry.partie,
                        "theme": entry.theme,
                        "titre": entry.titre,
                        "chapitre": entry.chapitre,
                        "section": entry.section,
                        "sous_section": entry.sous_section,
                        "reference": case.reference,
                        "description": case.description,
                        "source_page": case.source_page,
                    }
                )
                count += 1
        chapter_counts.append(
            {
                "chapter_number": chapter_number,
                "chapter_file": cache_path.stem + ".pdf",
                "jurisprudence_count": count,
                "status": "Traité",
            }
        )
    return rows, chapter_counts


def assemble_case_law_docx(cache_dir, output_docx_path, title=None):
    cache_paths = _sorted_cache_paths(cache_dir)
    if not cache_paths:
        raise ValueError("Aucun résultat de chapitre n'a été trouvé.")

    document = Document()
    _configure_doc_styles(document)
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "Répertoire des jurisprudences en encadrés"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    _add_page_number(section.footer.paragraphs[0])

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(24)
    title_paragraph.paragraph_format.space_after = Pt(6)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(
        title or "Répertoire des jurisprudences figurant dans les encadrés"
    )
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    subtitle_run = subtitle.add_run(
        "Classement conforme au plan du manuel - extraction visuelle par chapitre"
    )
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(10.5)
    subtitle_run.font.color.rgb = RGBColor(90, 90, 90)

    last_headings = [None] * 6
    heading_fields = [
        "partie", "theme", "titre", "chapitre", "section", "sous_section"
    ]
    case_count = 0
    for cache_path in cache_paths:
        extraction = _read_cache(cache_path)
        for entry in extraction.entries:
            values = [getattr(entry, field).strip() for field in heading_fields]
            for index, value in enumerate(values):
                if value and value != last_headings[index]:
                    document.add_heading(value, level=index + 1)
                    last_headings[index] = value
                    for deeper in range(index + 1, 6):
                        last_headings[deeper] = None
            for case in entry.jurisprudences:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(case.formulation.strip())
                case_count += 1

    if case_count == 0:
        document.add_paragraph("Aucune jurisprudence en encadré n'a été relevée.")

    output_path = Path(output_docx_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def assemble_case_law_xlsx(cache_dir, output_xlsx_path):
    rows, chapter_counts = _flatten_case_law_rows(cache_dir)
    if not chapter_counts:
        raise ValueError("Aucun résultat de chapitre n'a été trouvé.")

    workbook = openpyxl.Workbook()
    results_sheet = workbook.active
    results_sheet.title = "Jurisprudences"
    summary_sheet = workbook.create_sheet("Synthèse")

    result_headers = [
        "N° chapitre",
        "Fichier chapitre",
        "PDF source",
        "Partie",
        "Thème",
        "Titre",
        "Chapitre",
        "Section",
        "Sous-section",
        "Référence juridique",
        "Description synthétique",
        "Page source",
    ]
    result_keys = [
        "chapter_number",
        "chapter_file",
        "source_pdf",
        "partie",
        "theme",
        "titre",
        "chapitre",
        "section",
        "sous_section",
        "reference",
        "description",
        "source_page",
    ]
    results_sheet.append(result_headers)
    for row in rows:
        results_sheet.append([row[key] for key in result_keys])

    summary_headers = ["N° chapitre", "Fichier chapitre", "Jurisprudences", "Statut"]
    summary_sheet.append(summary_headers)
    for row in chapter_counts:
        summary_sheet.append(
            [
                row["chapter_number"],
                row["chapter_file"],
                row["jurisprudence_count"],
                row["status"],
            ]
        )
    total_row = summary_sheet.max_row + 1
    summary_sheet.cell(total_row, 2, "TOTAL")
    summary_sheet.cell(
        total_row,
        3,
        f"=SUM(C2:C{total_row - 1})",
    )

    header_fill = PatternFill("solid", fgColor="1F4D78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10)
    for sheet in (results_sheet, summary_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_cells in sheet.iter_rows(min_row=2):
            for cell in row_cells:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    result_widths = {
        "A": 12,
        "B": 20,
        "C": 42,
        "D": 34,
        "E": 34,
        "F": 34,
        "G": 42,
        "H": 42,
        "I": 36,
        "J": 42,
        "K": 70,
        "L": 12,
    }
    for column, width in result_widths.items():
        results_sheet.column_dimensions[column].width = width
    for column, width in {"A": 12, "B": 22, "C": 18, "D": 14}.items():
        summary_sheet.column_dimensions[column].width = width
    results_sheet.row_dimensions[1].height = 30
    summary_sheet.row_dimensions[1].height = 30

    if rows:
        result_table = ExcelTable(
            displayName="JurisprudenceResults",
            ref=f"A1:L{results_sheet.max_row}",
        )
        result_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        results_sheet.add_table(result_table)
    summary_table = ExcelTable(
        displayName="ChapterSummary",
        ref=f"A1:D{total_row - 1}",
    )
    summary_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    summary_sheet.add_table(summary_table)
    summary_sheet.cell(total_row, 2).font = Font(name="Calibri", bold=True)
    summary_sheet.cell(total_row, 3).font = Font(name="Calibri", bold=True)

    output_path = Path(output_xlsx_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def assemble_case_law_outputs(
    cache_dir,
    output_docx_path,
    output_xlsx_path=None,
    title=None,
):
    docx_path = assemble_case_law_docx(cache_dir, output_docx_path, title)
    xlsx_path = assemble_case_law_xlsx(
        cache_dir,
        output_xlsx_path or Path(output_docx_path).with_suffix(".xlsx"),
    )
    return docx_path, xlsx_path


async def _process_all(
    pdf_paths,
    example_image_path,
    cache_dir,
    settings,
    prompt,
    concurrency,
    retries,
    force,
):
    semaphore = asyncio.Semaphore(concurrency)

    async def process(pdf_path):
        cache_path = _cache_path(cache_dir, pdf_path)
        if cache_path.is_file() and not force:
            logging.info("Déjà traité : %s", pdf_path.name)
            return pdf_path, "cached", None
        try:
            async with semaphore:
                logging.info("Analyse visuelle : %s", pdf_path.name)
                extraction = await _extract_one(
                    settings, pdf_path, example_image_path, prompt, retries
                )
            _write_cache(cache_path, pdf_path, extraction)
            return pdf_path, "complete", None
        except Exception as exc:
            return pdf_path, "error", str(exc)

    results = []
    tasks = [asyncio.create_task(process(path)) for path in pdf_paths]
    for done, task in enumerate(asyncio.as_completed(tasks), start=1):
        result = await task
        results.append(result)
        logging.info("Progression %s/%s : %s", done, len(tasks), result[0].name)
    return results


def extract_boxed_case_law_chapters(
    chapters_dir,
    example_image_path,
    output_docx_path,
    output_xlsx_path=None,
    cache_dir=None,
    open_ai_key=None,
    model=None,
    concurrency=2,
    retries=3,
    env_path=".env",
    prompt=DEFAULT_PROMPT,
    force=False,
    from_chapter=None,
    to_chapter=None,
):
    chapters_path = Path(chapters_dir).resolve()
    example_path = Path(example_image_path).resolve()
    if not chapters_path.is_dir():
        raise ValueError(f"Dossier de chapitres introuvable : {chapters_path}")
    if not example_path.is_file():
        raise ValueError(f"Image d'exemple introuvable : {example_path}")
    if concurrency < 1:
        raise ValueError("concurrency doit être supérieur ou égal à 1.")

    pdf_paths = sorted(chapters_path.glob("*.pdf"), key=_chapter_number)
    if from_chapter is not None:
        pdf_paths = [p for p in pdf_paths if _chapter_number(p) >= from_chapter]
    if to_chapter is not None:
        pdf_paths = [p for p in pdf_paths if _chapter_number(p) <= to_chapter]
    if not pdf_paths:
        raise ValueError("Aucun PDF de chapitre ne correspond à la sélection.")

    output_path = Path(output_docx_path).resolve()
    cache_path = (
        Path(cache_dir).resolve()
        if cache_dir
        else output_path.parent / f"{output_path.stem}_chapters"
    )
    settings = _resolve_openai_settings(
        open_ai_key=open_ai_key,
        model=model,
        env_path=env_path,
    )
    if not settings["api_key"]:
        raise ValueError("La clé OpenAI est absente.")

    results = asyncio.run(
        _process_all(
            pdf_paths,
            example_path,
            cache_path,
            settings,
            prompt,
            concurrency,
            retries,
            force,
        )
    )
    errors = [(path.name, error) for path, status, error in results if status == "error"]
    if errors:
        details = " | ".join(f"{name}: {error}" for name, error in errors)
        raise RuntimeError(
            f"{len(errors)} chapitre(s) en erreur. Les réussites sont conservées : {details}"
        )
    return assemble_case_law_outputs(
        cache_path,
        output_path,
        output_xlsx_path=output_xlsx_path,
    )
