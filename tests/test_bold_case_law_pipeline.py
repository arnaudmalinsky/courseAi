import json
import tempfile
import unittest
from pathlib import Path

import fitz
import openpyxl
from docx import Document

from text_processing.bold_case_law_pipeline import (
    BoldChapterExtraction,
    _build_input,
    _normalize_heading_text,
    assemble_bold_case_law_outputs,
    build_chapter_slices,
    split_pdf_by_detected_chapters,
)


class BoldCaseLawPipelineTests(unittest.TestCase):
    def _make_pdf(self, path):
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text((70, 80), "PARTIE I - DROIT PUBLIC", fontsize=18)
        page.insert_text((70, 120), "THEME I - SOURCES", fontsize=16)
        page.insert_text((70, 250), "CHAPITRE I - PREMIER", fontsize=16)
        page.insert_text((70, 300), "Contenu du premier chapitre", fontsize=10)
        page.insert_text((70, 500), "CHAPITRE II - SECOND", fontsize=16)
        page.insert_text((70, 550), "Contenu du second chapitre", fontsize=10)
        toc = doc.new_page(width=595, height=842)
        toc.insert_text((70, 80), "Chapitre I .......... 1", fontsize=12)
        toc.insert_text((70, 110), "Chapitre II ......... 2", fontsize=12)
        doc.save(path)
        doc.close()

    def test_heading_normalization(self):
        self.assertEqual(
            "CHAPITRE II — LES TITRES",
            _normalize_heading_text("C HAPITRE II — L ES TITRES"),
        )

    def test_split_handles_two_chapters_on_same_page(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.pdf"
            self._make_pdf(source)
            with fitz.open(source) as doc:
                slices, end = build_chapter_slices(doc)
            self.assertEqual(2, len(slices))
            self.assertEqual(slices[0].end_page_index, slices[1].start_page_index)
            self.assertAlmostEqual(slices[0].end_y, slices[1].start_y)
            manifest = split_pdf_by_detected_chapters(source, root / "chapters")
            payload = json.loads(manifest.read_text(encoding="utf-8"))
            self.assertEqual("complete", payload["coverage_status"])
            self.assertEqual(2, payload["chapter_count"])
            for chapter in payload["chapters"]:
                self.assertTrue(Path(chapter["chapter_pdf"]).is_file())
            first = fitz.open(payload["chapters"][0]["chapter_pdf"])
            second = fitz.open(payload["chapters"][1]["chapter_pdf"])
            self.assertNotIn("CHAPITRE II", first[0].get_text("text"))
            self.assertEqual("", first[0].get_text("text"))
            self.assertEqual("", second[0].get_text("text"))
            first.close()
            second.close()

    def test_input_has_no_example_image(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf = Path(temp_dir) / "chapter.pdf"
            self._make_pdf(pdf)
            record = {
                "chapter_pdf": str(pdf),
                "source_start_page": 1,
                "source_end_page": 1,
            }
            content = _build_input(record, "PROMPT")[0]["content"]
            self.assertEqual(["input_text", "input_file"], [item["type"] for item in content])

    def test_grouped_references_share_one_output_row(self):
        payload = {
            "entries": [
                {
                    "partie": "PARTIE I",
                    "theme": "THÈME I",
                    "titre": "",
                    "chapitre": "CHAPITRE I",
                    "section": "SECTION I",
                    "sous_section": "",
                    "groupes": [
                        {
                            "references": ["CE, 1950, A", "CE, 1951, B"],
                            "description": "consacrent la même règle.",
                            "source_page": 12,
                        }
                    ],
                }
            ]
        }
        extraction = BoldChapterExtraction.model_validate(payload)
        self.assertEqual("CE, 1950, A ; CE, 1951, B", extraction.entries[0].groupes[0].reference_text)
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            cache = root / "cache"
            cache.mkdir()
            (cache / "chapitre_001.json").write_text(
                json.dumps(
                    {
                        "chapter": {
                            "order": 1,
                            "chapter_pdf": str(root / "chapter.pdf"),
                            "source_pdf": str(root / "source.pdf"),
                        },
                        "result": payload,
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            docx_path, xlsx_path = assemble_bold_case_law_outputs(
                cache, root / "result.docx"
            )
            sheet = openpyxl.load_workbook(xlsx_path)["Jurisprudences"]
            self.assertEqual(2, sheet.max_row)
            self.assertEqual("CE, 1950, A ; CE, 1951, B", sheet["J2"].value)
            self.assertEqual(2, sheet["K2"].value)
            document = Document(docx_path)
            self.assertTrue(any("CE, 1950, A ; CE, 1951, B" in p.text for p in document.paragraphs))


if __name__ == "__main__":
    unittest.main()
