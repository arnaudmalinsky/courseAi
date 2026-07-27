import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from docx import Document

from text_processing.boxed_case_law_pipeline import (
    ChapterExtraction,
    _build_input,
    _parse_extraction,
    assemble_case_law_docx,
    assemble_case_law_outputs,
    extract_boxed_case_law_chapters,
)


class BoxedCaseLawPipelineTests(unittest.TestCase):
    def test_build_input_contains_prompt_image_and_pdf_without_overlap(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            image = root / "example.jpeg"
            pdf = root / "chapitre_1.pdf"
            image.write_bytes(b"jpeg")
            pdf.write_bytes(b"pdf")
            request = _build_input(pdf, image, "PROMPT")
            content = request[0]["content"]
            self.assertEqual(["input_text", "input_image", "input_file"], [x["type"] for x in content])
            self.assertEqual("high", content[1]["detail"])
            self.assertEqual("high", content[2]["detail"])

    def test_parse_and_assemble_preserve_plan_and_formulation(self):
        payload = {
            "entries": [
                {
                    "partie": "PARTIE I - SOURCES",
                    "theme": "THÈME I - PRINCIPES",
                    "titre": "",
                    "chapitre": "CHAPITRE I - CONTRÔLE",
                    "section": "SECTION 1 - VALIDATIONS",
                    "sous_section": "",
                    "jurisprudences": [
                        {
                            "reference": "CE, 28 juillet 2000, Tête",
                            "description": "contrôle les lois de validation.",
                            "source_page": 2,
                        }
                    ],
                }
            ]
        }
        extraction = _parse_extraction(json.dumps(payload, ensure_ascii=False))
        self.assertIsInstance(extraction, ChapterExtraction)
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            cache = root / "cache"
            cache.mkdir()
            (cache / "chapitre_1.json").write_text(
                json.dumps({"source_pdf": "chapitre_1.pdf", "result": payload}, ensure_ascii=False),
                encoding="utf-8",
            )
            output = assemble_case_law_docx(cache, root / "result.docx")
            document = Document(output)
            texts = [p.text for p in document.paragraphs]
            self.assertIn("PARTIE I - SOURCES", texts)
            self.assertIn("CHAPITRE I - CONTRÔLE", texts)
            self.assertTrue(any("Tête : contrôle" in text for text in texts))

    def test_pipeline_caches_and_resumes(self):
        calls = []

        async def fake_call(settings, prompt, text_format=None):
            calls.append((settings, prompt, text_format))
            return '{"entries": []}'

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            chapters = root / "chapters"
            chapters.mkdir()
            (chapters / "chapitre_1.pdf").write_bytes(b"pdf")
            image = root / "example.jpeg"
            image.write_bytes(b"jpeg")
            output = root / "result.docx"
            with patch(
                "text_processing.boxed_case_law_pipeline._call_openai_responses",
                side_effect=fake_call,
            ):
                extract_boxed_case_law_chapters(
                    chapters, image, output, open_ai_key="test-key"
                )
                extract_boxed_case_law_chapters(
                    chapters, image, output, open_ai_key="test-key"
                )
            self.assertEqual(1, len(calls))
            self.assertTrue(output.is_file())
            self.assertTrue(output.with_suffix(".xlsx").is_file())

    def test_combined_assembly_creates_docx_and_xlsx(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            cache = root / "cache"
            cache.mkdir()
            empty = {"source_pdf": "chapitre_1.pdf", "result": {"entries": []}}
            result = {
                "source_pdf": "chapitre_2.pdf",
                "result": {
                    "entries": [
                        {
                            "partie": "PARTIE I",
                            "theme": "THÈME I",
                            "titre": "",
                            "chapitre": "CHAPITRE I",
                            "section": "SECTION I",
                            "sous_section": "",
                            "jurisprudences": [
                                {
                                    "reference": "CE, 13 décembre 1889, Cadot",
                                    "description": "consacre le Conseil d'État comme juge de droit commun.",
                                    "source_page": 2,
                                }
                            ],
                        }
                    ]
                },
            }
            (cache / "chapitre_1.json").write_text(
                json.dumps(empty, ensure_ascii=False), encoding="utf-8"
            )
            (cache / "chapitre_2.json").write_text(
                json.dumps(result, ensure_ascii=False), encoding="utf-8"
            )
            docx_path, xlsx_path = assemble_case_law_outputs(
                cache, root / "result.docx"
            )
            self.assertTrue(docx_path.is_file())
            self.assertTrue(xlsx_path.is_file())
            workbook = openpyxl.load_workbook(xlsx_path, data_only=False)
            self.assertEqual(["Jurisprudences", "Synthèse"], workbook.sheetnames)
            self.assertEqual(
                "CE, 13 décembre 1889, Cadot",
                workbook["Jurisprudences"]["J2"].value,
            )
            self.assertIn(
                "juge de droit commun",
                workbook["Jurisprudences"]["K2"].value,
            )
            self.assertEqual("=SUM(C2:C3)", workbook["Synthèse"]["C4"].value)

    def test_old_formulation_cache_is_split_for_backward_compatibility(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            cache = root / "cache"
            cache.mkdir()
            payload = {
                "source_pdf": "chapitre_2.pdf",
                "result": {
                    "entries": [
                        {
                            "partie": "",
                            "theme": "",
                            "titre": "",
                            "chapitre": "CHAPITRE I",
                            "section": "",
                            "sous_section": "",
                            "jurisprudences": [
                                {
                                    "formulation": "CE, 13 décembre 1889, Cadot : consacre le Conseil d'État comme juge de droit commun.",
                                    "source_page": 2,
                                }
                            ],
                        }
                    ]
                },
            }
            (cache / "chapitre_2.json").write_text(
                json.dumps(payload, ensure_ascii=False), encoding="utf-8"
            )
            _, xlsx_path = assemble_case_law_outputs(cache, root / "result.docx")
            sheet = openpyxl.load_workbook(xlsx_path)["Jurisprudences"]
            self.assertEqual("CE, 13 décembre 1889, Cadot", sheet["J2"].value)
            self.assertIn("juge de droit commun", sheet["K2"].value)


if __name__ == "__main__":
    unittest.main()
