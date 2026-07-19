import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from docx import Document

from text_processing.chapter_pipeline import (
    assemble_chapter_summaries,
    split_docx_by_chapter,
    summarize_chapter_manifest,
    validate_summary,
)


class ChapterPipelineTests(unittest.TestCase):
    def test_split_preserves_every_block_and_creates_manifest(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "course.docx"
            chapters_dir = root / "chapters"

            document = Document()
            document.add_paragraph("Avant-propos à conserver.")
            document.add_heading("PARTIE I - SOURCES", level=1)
            document.add_heading("THÈME I - PRINCIPES", level=2)
            document.add_heading("CHAPITRE I - PRINCIPES", level=3)
            document.add_heading("Section 1 - Règle", level=4)
            document.add_paragraph("Texte du premier chapitre.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Article"
            table.cell(0, 1).text = "Contenu"
            document.add_heading("PARTIE II - APPLICATIONS", level=1)
            document.add_heading("THÈME II - CAS", level=2)
            document.add_heading("CHAPITRE II - APPLICATION", level=3)
            document.add_paragraph("Texte du second chapitre.")
            document.save(source_path)

            manifest_path = split_docx_by_chapter(source_path, chapters_dir)

            workbook = openpyxl.load_workbook(manifest_path)
            worksheet = workbook["Chapters"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            rows = list(worksheet.iter_rows(min_row=2, values_only=True))
            self.assertEqual(3, len(rows))
            self.assertEqual("preamble", worksheet.cell(2, headers["chapter_id"]).value)
            self.assertEqual(
                "complete", worksheet.cell(2, headers["completeness_status"]).value
            )
            self.assertTrue(
                all(
                    Path(worksheet.cell(row, headers["source_docx"]).value).is_file()
                    for row in range(2, worksheet.max_row + 1)
                )
            )
            self.assertEqual(
                ["Préambule", "CHAPITRE I - PRINCIPES", "CHAPITRE II - APPLICATION"],
                [
                    worksheet.cell(row, headers["title"]).value
                    for row in range(2, worksheet.max_row + 1)
                ],
            )
            self.assertEqual(
                ("PARTIE I - SOURCES", "THÈME I - PRINCIPES"),
                (
                    worksheet.cell(3, headers["partie"]).value,
                    worksheet.cell(3, headers["theme"]).value,
                ),
            )
            self.assertEqual(
                ("PARTIE II - APPLICATIONS", "THÈME II - CAS"),
                (
                    worksheet.cell(4, headers["partie"]).value,
                    worksheet.cell(4, headers["theme"]).value,
                ),
            )
            first_chapter = Document(
                worksheet.cell(3, headers["source_docx"]).value
            )
            self.assertNotIn(
                "PARTIE II - APPLICATIONS",
                [paragraph.text for paragraph in first_chapter.paragraphs],
            )

    def test_split_rejects_document_without_chapter(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "course.docx"
            document = Document()
            document.add_heading("Introduction", level=1)
            document.add_paragraph("Texte.")
            document.save(source_path)

            with self.assertRaisesRegex(ValueError, "Aucun titre"):
                split_docx_by_chapter(source_path, root / "chapters")

    def test_table_of_contents_entries_are_excluded(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "course.docx"
            document = Document()
            document.add_heading("CHAPITRE I - COURS", level=3)
            document.add_paragraph("Contenu réel.")
            document.add_heading(
                "Chapitre I - Cours ........................................ 12",
                level=3,
            )
            document.save(source_path)

            manifest_path = split_docx_by_chapter(source_path, root / "chapters")
            worksheet = openpyxl.load_workbook(manifest_path)["Chapters"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            self.assertEqual(
                "table_of_contents",
                worksheet.cell(3, headers["section_type"]).value,
            )
            self.assertEqual("excluded", worksheet.cell(3, headers["status"]).value)

    def test_table_of_contents_entry_with_page_number_only_is_excluded(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "course.docx"
            document = Document()
            document.add_heading("CHAPITRE II : LE COURS", level=3)
            document.add_paragraph("Contenu réel.")
            document.add_heading(
                "Chapitre II : Le cours 60",
                level=3,
            )
            document.save(source_path)

            manifest_path = split_docx_by_chapter(source_path, root / "chapters")
            worksheet = openpyxl.load_workbook(manifest_path)["Chapters"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            self.assertEqual(
                "table_of_contents",
                worksheet.cell(3, headers["section_type"]).value,
            )
            self.assertEqual("excluded", worksheet.cell(3, headers["status"]).value)

    def test_summary_validation_checks_headings_and_reduction(self):
        source = (
            "### CHAPITRE I\n\n#### Section A\n\n"
            + "Une règle juridique essentielle et ses conditions. " * 30
        )
        summary = (
            "### CHAPITRE I\n\n#### Section A\n\n"
            "La **règle juridique** s’applique sous conditions."
        )
        errors, _, _, reduction = validate_summary(
            source,
            [(3, "CHAPITRE I"), (4, "Section A")],
            False,
            summary,
        )
        self.assertEqual([], errors)
        self.assertGreaterEqual(reduction, 0.5)

    def test_summary_processing_selects_order_range_saves_result_and_resumes(self):
        calls = []

        async def fake_call(settings, prompt):
            calls.append((settings, prompt))
            return "### CHAPITRE II\n\nSynthèse fidèle."

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "course.docx"
            document = Document()
            document.add_heading("CHAPITRE I", level=3)
            document.add_paragraph("Première règle juridique essentielle.")
            document.add_heading("CHAPITRE II", level=3)
            document.add_paragraph("Deuxième règle juridique essentielle.")
            document.save(source_path)
            manifest_path = split_docx_by_chapter(source_path, root / "chapters")

            with patch(
                "text_processing.chapter_pipeline._call_openai_responses",
                side_effect=fake_call,
            ):
                with self.assertLogs(level="INFO") as captured_logs:
                    summarize_chapter_manifest(
                        manifest_path,
                        "test-key",
                        concurrency=2,
                        from_order=1,
                        to_order=1,
                    )
                summarize_chapter_manifest(
                    manifest_path,
                    "test-key",
                    concurrency=2,
                    from_order=1,
                    to_order=1,
                )
            log_output = "\n".join(captured_logs.output)
            self.assertIn("Progression 1/1 (100.0%)", log_output)
            self.assertIn("Traitement terminé", log_output)

            workbook = openpyxl.load_workbook(manifest_path)
            worksheet = workbook["Chapters"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            self.assertEqual("pending", worksheet.cell(2, headers["status"]).value)
            self.assertEqual("complete", worksheet.cell(3, headers["status"]).value)
            self.assertEqual("OK", worksheet.cell(3, headers["validation"]).value)
            self.assertEqual(1, len(calls))

            with self.assertRaisesRegex(ValueError, "--from-order"):
                summarize_chapter_manifest(
                    manifest_path,
                    "test-key",
                    from_order=2,
                    to_order=1,
                )

    def test_assembly_keeps_upper_titles_and_ignores_validation_warnings(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_path = root / "course.docx"
            document = Document()
            document.add_heading("PARTIE I", level=1)
            document.add_heading("THÈME I", level=2)
            document.add_heading("CHAPITRE I", level=3)
            document.add_paragraph("Texte source suffisamment long.")
            document.save(source_path)
            manifest_path = split_docx_by_chapter(source_path, root / "chapters")

            workbook = openpyxl.load_workbook(manifest_path)
            worksheet = workbook["Chapters"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            worksheet.cell(2, headers["status"], "invalid")
            worksheet.cell(
                2,
                headers["summary"],
                "### CHAPITRE I\n\nLa **règle** est synthétisée.",
            )
            worksheet.cell(2, headers["source_length"], 100)
            worksheet.cell(2, headers["summary_length"], 60)
            worksheet.cell(2, headers["validation"], "WARNING: test")
            workbook.save(manifest_path)

            output_path = assemble_chapter_summaries(
                manifest_path, root / "result.docx"
            )
            result = Document(output_path)
            self.assertEqual(
                [
                    ("Heading 1", "PARTIE I"),
                    ("Heading 2", "THÈME I"),
                    ("Heading 3", "CHAPITRE I"),
                ],
                [
                    (paragraph.style.name, paragraph.text)
                    for paragraph in result.paragraphs[:3]
                ],
            )
            self.assertEqual("La règle est synthétisée.", result.paragraphs[3].text)
            self.assertTrue(result.paragraphs[3].runs[1].bold)


if __name__ == "__main__":
    unittest.main()
